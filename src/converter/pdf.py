from __future__ import annotations

import re
import shutil
from collections import Counter
from itertools import pairwise
from pathlib import Path
from statistics import median

from .config import environment, resource_root
from .models import Block, Document, file_hash
from .pdf_layout import add_page_context, column_bands, merge_paragraph_lines, word_style
from .scan_layout import remaining_figures, ruled_tables

PDF_ADAPTER_VERSION = 4


def configure_tesseract():
    import pytesseract

    bundled = resource_root() / "assets/tesseract"
    explicit = environment("TESSERACT")
    executable = Path(explicit) if explicit else bundled / "tesseract.exe"
    if executable.exists():
        pytesseract.pytesseract.tesseract_cmd = str(executable)
    elif shutil.which("tesseract"):
        pytesseract.pytesseract.tesseract_cmd = shutil.which("tesseract")
    else:
        raise ValueError("Không tìm thấy Tesseract. Cần bộ OCR đi kèm để đọc PDF scan.")
    tessdata = bundled / "tessdata"
    return f'--tessdata-dir "{tessdata}"' if tessdata.exists() else ""


def render_page(source: Path, page_index: int, scale: float = 2.0):
    import pypdfium2 as pdfium

    with pdfium.PdfDocument(source) as pdf:
        page = pdf[page_index]
        try:
            bitmap = page.render(scale=scale)
            try:
                return bitmap.to_pil().copy()
            finally:
                bitmap.close()
        finally:
            page.close()


def ocr_image(image, language: str, config: str, psm=3) -> list[dict]:
    import pytesseract

    installed = set(pytesseract.get_languages(config=config))
    missing = set(language.split("+")) - installed
    if missing:
        raise ValueError("Thiếu OCR language packs: " + ", ".join(sorted(missing)))
    data = pytesseract.image_to_data(
        image, lang=language, config=config + f" --psm {psm}", output_type=pytesseract.Output.DICT
    )
    grouped = {}
    for index, text in enumerate(data["text"]):
        if not text.strip():
            continue
        key = tuple(data[name][index] for name in ["block_num", "par_num", "line_num"])
        grouped.setdefault(key, []).append(index)
    result = []
    for indices in grouped.values():
        left = min(data["left"][i] for i in indices)
        top = min(data["top"][i] for i in indices)
        right = max(data["left"][i] + data["width"][i] for i in indices)
        bottom = max(data["top"][i] + data["height"][i] for i in indices)
        result.append(
            {
                "text": join_words([data["text"][i] for i in indices]),
                "bbox": [left, top, right, bottom],
                "confidence": sum(float(data["conf"][i]) for i in indices) / len(indices),
                "minimum_confidence": min(float(data["conf"][i]) for i in indices),
                "paragraph": [data["block_num"][indices[0]], data["par_num"][indices[0]]],
                "text_height": median(data["height"][i] for i in indices),
            }
        )
    return result


def join_words(words: list[str]) -> str:
    result = ""
    han = re.compile(r"[\u3400-\u9fff\uf900-\ufaff]")
    for word in words:
        if result and word and not (han.fullmatch(result[-1]) and han.fullmatch(word[0])):
            result += " "
        result += word
    return result


def ocr_confidence(lines: list[dict]) -> float:
    weights = [max(1, len(line["text"])) for line in lines]
    return sum(line["confidence"] * weight for line, weight in zip(lines, weights, strict=True)) / max(1, sum(weights))


def refine_ocr_lines(image, lines: list[dict], language: str, config: str, *, fragmented=False) -> list[dict]:
    """Retry weak physical text lines, preserving their page/paragraph anchors."""
    result = []
    for line in lines:
        minimum = line.get("minimum_confidence", line["confidence"])
        if minimum >= (85 if fragmented else 65) or len(line["text"].strip()) < 6:
            result.append(line)
            continue
        left, top, right, bottom = line["bbox"]
        height = line.get("text_height", bottom - top)
        if right - left < height * 3:
            result.append(line)
            continue
        pad = max(12, round(height / 2))
        crop = image.crop((max(0, left - pad), max(0, top - pad),
                           min(image.width, right + pad), min(image.height, bottom + pad)))
        try:
            retry = ocr_image(crop, language, config, 7)
            candidates = retry if len(retry) == 1 else []
            if fragmented:
                from PIL import ImageFilter
                emphasized = crop.filter(ImageFilter.MinFilter(3))
                try:
                    retry = ocr_image(emphasized, language, config, 7)
                    if len(retry) == 1:
                        candidates += retry
                finally:
                    emphasized.close()
        finally:
            crop.close()
        # Preserve existing numbers/product codes when reconsidering uncertain
        # line segmentation. Higher confidence alone is not evidence to change them.
        protected = r"(?<![A-Za-z0-9_])[A-Za-z]+[-_][A-Za-z0-9_-]*\d[A-Za-z0-9_-]*|\d+(?:[.,:/-]\d+)*"
        gain, minimum_gain = (1, 2) if fragmented else (5, 10)
        eligible = [candidate for candidate in candidates
                    if candidate["confidence"] >= line["confidence"] + gain
                    and candidate.get("minimum_confidence", candidate["confidence"]) >= minimum + minimum_gain
                    and len(candidate["text"].strip()) >= len(line["text"].strip()) * .85
                    and len(candidate["text"].strip()) <= len(line["text"].strip()) * 1.25
                    and Counter(re.findall(protected, candidate["text"])) == Counter(re.findall(protected, line["text"]))]
        if eligible:
            retry = max(eligible, key=lambda candidate:candidate["confidence"])
            result.append(line | {key:retry[key] for key in
                                  ["text", "confidence", "minimum_confidence"] if key in retry})
        else:
            result.append(line)
    return result


def read_pdf(
    source: Path, assets_dir: Path, language="chi_sim+chi_tra+eng", progress=None, control=None
) -> Document:
    import pdfplumber

    doc = Document(str(source), file_hash(source), "pdf", adapter_version=PDF_ADAPTER_VERSION)
    assets_dir.mkdir(parents=True, exist_ok=True)
    doc.warnings.append("PDF → DOCX dựng lại nội dung; bố cục/font/số trang có thể khác bản gốc.")
    config = None
    with pdfplumber.open(source) as pdf:
        if getattr(pdf.doc, "encryption", None):
            raise ValueError("Chưa hỗ trợ PDF được mã hóa.")
        for number, page in enumerate(pdf.pages):
            if control:
                control()
            if progress:
                progress(f"Đọc PDF: trang {number + 1}/{len(pdf.pages)}")
            section = Block("section", meta={
                "title": f"Page {number + 1}", "page": number,
                "width": float(page.width), "height": float(page.height),
            })
            words = page.extract_words(extra_attrs=["size", "fontname", "non_stroking_color"])
            baseline = (
                median([float(char.get("size", 12)) for char in page.chars]) if page.chars else 12
            )
            ordered: list[tuple[float, float, Block]] = []
            tables = page.find_tables() if words else []
            for table in tables:
                block = Block("table", meta={"bbox": list(table.bbox)})
                for row_index, row in enumerate(table.extract()):
                    rb = Block("row")
                    for column, text in enumerate(row):
                        bounds = table.rows[row_index].cells[column] or table.bbox
                        cell = Block("cell", meta={"bbox": list(bounds)})
                        if text:
                            segment = doc.add_segment(
                                text,
                                {
                                    "page": number,
                                    "bbox": list(bounds),
                                    "row": row_index,
                                    "column": column,
                                },
                                context=text,
                                style=word_style([
                                    word for word in words
                                    if bounds[0] <= (word["x0"] + word["x1"]) / 2 <= bounds[2]
                                    and bounds[1] <= (word["top"] + word["bottom"]) / 2 <= bounds[3]
                                ]),
                            )
                            cell.segments.append(segment.id)
                        rb.children.append(cell)
                    block.children.append(rb)
                ordered.append((table.bbox[1], table.bbox[0], block))

            if words:
                # Use layout-aware lines outside table regions; avoid translating tables twice.
                normal = [
                    word
                    for word in words
                    if not any(
                        table.bbox[0] <= (word["x0"] + word["x1"]) / 2 <= table.bbox[2]
                        and table.bbox[1] <= (word["top"] + word["bottom"]) / 2 <= table.bbox[3]
                        for table in tables
                    )
                ]
                lines = []
                for word in sorted(normal, key=lambda item: (round(item["top"] / 3), item["x0"])):
                    if (
                        lines
                        and abs(lines[-1][0]["top"] - word["top"]) < 3
                        and word["x0"] - lines[-1][-1]["x1"] < 30
                    ):
                        lines[-1].append(word)
                    else:
                        lines.append([word])
                for line in lines:
                    text = join_words([word["text"] for word in line])
                    bbox = [
                        min(w["x0"] for w in line),
                        min(w["top"] for w in line),
                        max(w["x1"] for w in line),
                        max(w["bottom"] for w in line),
                    ]
                    segment = doc.add_segment(
                        text, {"page": number, "bbox": bbox}, context=text, style=word_style(line)
                    )
                    size = max(float(word.get("size", baseline)) for word in line)
                    heading = size >= baseline * 1.2 and len(text) <= 200
                    ordered.append(
                        (
                            bbox[1],
                            bbox[0],
                            Block(
                                "heading" if heading else "paragraph",
                                [segment.id],
                                meta={"bbox": bbox} | (
                                    {"level": 1 if size >= baseline * 1.5 else 2} if heading else {}
                                ),
                            ),
                        )
                    )
            else:
                if config is None:
                    config = configure_tesseract()
                image = render_page(source, number, scale=3)
                try:
                    recognized = ocr_image(image, language, config)
                    fragmented = False
                    narrow = sum(item["bbox"][2] - item["bbox"][0]
                                 < (item["bbox"][3] - item["bbox"][1]) * 2.5
                                 for item in recognized)
                    if (3 <= len(recognized) <= 12 and narrow >= len(recognized) / 2
                            and any(item["confidence"] < 50 for item in recognized)):
                        sparse = ocr_image(image, language, config, 11)
                        if sparse and ocr_confidence(sparse) > ocr_confidence(recognized) + 5:
                            recognized = sparse
                            fragmented = True
                    grids = ruled_tables(image)
                    table_boxes = [grid["bbox"] for grid in grids]
                    text_height = median([item.get("text_height", item["bbox"][3] - item["bbox"][1])
                                          for item in recognized]) if recognized else 36

                    def keep_region(pixel_box, suffix, number=number, image=image, ordered=ordered):
                        asset = f"page-{number + 1}-scan-{suffix}.png"
                        path = assets_dir / asset
                        crop = image.crop(tuple(pixel_box))
                        try:
                            crop.save(path)
                        finally:
                            crop.close()
                        doc.assets[asset] = str(path)
                        bounds = [value / 3 for value in pixel_box]
                        ordered.append((bounds[1], bounds[0], Block("image", meta={
                            "asset": asset, "bbox": bounds, "scan": True,
                        })))

                    for index, grid in enumerate(grids):
                        cells, readable = [], True
                        for top, bottom in pairwise(grid["ys"]):
                            row = []
                            for left, right in pairwise(grid["xs"]):
                                crop = image.crop((left + 4, top + 4, right - 4, bottom - 4))
                                try:
                                    lines = ocr_image(crop, language, config, 6)
                                    if lines and min(line["confidence"] for line in lines) < 65:
                                        enlarged = crop.resize((crop.width * 2, crop.height * 2))
                                        try:
                                            retry = ocr_image(enlarged, language, config, 7)
                                        finally:
                                            enlarged.close()
                                        if retry and min(line["confidence"] for line in retry) > (
                                                min(line["confidence"] for line in lines)):
                                            for line in retry:
                                                if "text_height" in line:
                                                    line["text_height"] /= 2
                                            lines = retry
                                    text = "\n".join(line["text"] for line in lines)
                                    if not text:
                                        gray = crop.convert("L")
                                        try:
                                            if gray.getextrema()[0] < 150:
                                                readable = False
                                        finally:
                                            gray.close()
                                    if lines and min(line["confidence"] for line in lines) < 65:
                                        readable = False
                                finally:
                                    crop.close()
                                row.append((text, [left / 3, top / 3, right / 3, bottom / 3], lines))
                            cells.append(row)
                        if not readable:
                            keep_region(grid["bbox"], f"table-{index}")
                            doc.warnings.append(f"Trang {number + 1}: bảng scan chưa đọc đủ; giữ ảnh gốc.")
                            continue
                        block = Block("table", meta={"bbox": [v / 3 for v in grid["bbox"]]})
                        for row_index, row in enumerate(cells):
                            rb = Block("row")
                            for column, (text, bounds, lines) in enumerate(row):
                                cell = Block("cell", meta={"bbox": bounds})
                                if text:
                                    size = median([line.get("text_height", text_height)
                                                   for line in lines]) / 3
                                    segment = doc.add_segment(text, {
                                        "page": number, "bbox": bounds, "row": row_index,
                                        "column": column, "ocr_table": True,
                                    }, context=text, style={"font_size": round(max(6, size), 1)})
                                    if min(line.get("minimum_confidence", line["confidence"])
                                           for line in lines) < 65:
                                        segment.warning = "Ô bảng scan: kiểm tra text OCR."
                                    cell.segments.append(segment.id)
                                rb.children.append(cell)
                            block.children.append(rb)
                        ordered.append((block.meta["bbox"][1], block.meta["bbox"][0], block))
                    recognized = refine_ocr_lines(image, recognized, language, config, fragmented=fragmented)
                    figures = remaining_figures(image, [item["bbox"] for item in recognized],
                                                table_boxes, text_height)
                    for index, bounds in enumerate(figures):
                        keep_region(bounds, f"figure-{index}")
                    if figures:
                        doc.warnings.append(f"Trang {number + 1}: giữ vùng ảnh scan; chữ trong ảnh giữ nguyên.")
                    recognized = [item for item in recognized if not any(
                        box[0] <= (item["bbox"][0] + item["bbox"][2]) / 2 <= box[2]
                        and box[1] <= (item["bbox"][1] + item["bbox"][3]) / 2 <= box[3]
                        for box in table_boxes
                    )]
                finally:
                    image.close()
                ocr_size = median([
                    item.get("text_height", item["bbox"][3] - item["bbox"][1]) / 3
                    for item in recognized
                ]) if recognized else 12
                if not recognized and not any(block.kind == "table" for _, _, block in ordered):
                    doc.warnings.append(
                        f"Trang {number + 1}: OCR không nhận được text; cần review trang nguồn."
                    )
                for item in recognized:
                    bbox = [value / 3 for value in item["bbox"]]
                    segment = doc.add_segment(
                        item["text"], {
                            "page": number, "bbox": bbox,
                            "ocr_paragraph": item["paragraph"],
                        }, context=item["text"], style={"font_size": round(max(6, item.get("text_height", ocr_size * 3) / 3), 1)}
                    )
                    confidence = item.get("minimum_confidence", item["confidence"])
                    if confidence < 65:
                        segment.warning = (
                            f"OCR confidence {confidence:.0f}%: kiểm tra text nguồn."
                        )
                    heading = segment.style["font_size"] > ocr_size * 1.3 and len(item["text"]) < 200
                    ordered.append((bbox[1], bbox[0], Block(
                        "heading" if heading else "paragraph", [segment.id],
                        meta={"bbox": bbox} | ({"level": 2} if heading else {}),
                    )))
                doc.warnings.append(f"Trang {number + 1}: dùng OCR; bảng và thứ tự đọc cần review.")
            # Keep embedded images; do not duplicate a scanned full-page image alongside its OCR.
            bitmap = None
            for index, item in enumerate(page.images):
                area = (item["x1"] - item["x0"]) * (item["bottom"] - item["top"])
                if not words and area > page.width * page.height * 0.8:
                    continue
                if bitmap is None:
                    bitmap = render_page(source, number, scale=2)
                bounds = [
                    max(0, item["x0"]),
                    max(0, item["top"]),
                    min(page.width, item["x1"]),
                    min(page.height, item["bottom"]),
                ]
                if bounds[2] <= bounds[0] or bounds[3] <= bounds[1]:
                    continue
                asset = f"page-{number + 1}-image-{index}.png"
                path = assets_dir / asset
                bitmap.crop(tuple(int(value * 2) for value in bounds)).save(path)
                doc.assets[asset] = str(path)
                ordered.append((bounds[1], bounds[0], Block(
                    "image", meta={"asset": asset, "bbox": bounds}
                )))
                # Text inside a raster image on a text-bearing PDF needs its own OCR pass.
                if words and area > 400:
                    if config is None:
                        config = configure_tesseract()
                    crop = bitmap.crop(tuple(int(value * 2) for value in bounds))
                    for item in ocr_image(crop, language, config):
                        bbox = [
                            bounds[0] + item["bbox"][0] / 2,
                            bounds[1] + item["bbox"][1] / 2,
                            bounds[0] + item["bbox"][2] / 2,
                            bounds[1] + item["bbox"][3] / 2,
                        ]
                        # Some PDFs embed the same image beneath an existing OCR text layer.
                        if any(
                            bbox[0] <= (w["x0"] + w["x1"]) / 2 <= bbox[2]
                            and bbox[1] <= (w["top"] + w["bottom"]) / 2 <= bbox[3]
                            for w in words
                        ):
                            continue
                        segment = doc.add_segment(
                            item["text"], {"page": number, "bbox": bbox}, context=item["text"]
                        )
                        segment.warning = "Text từ ảnh: kiểm tra OCR; ảnh gốc vẫn giữ nguyên."
                        ordered.append((bbox[1], bbox[0], Block(
                            "paragraph", [segment.id], meta={"bbox": bbox}
                        )))
            if bitmap:
                bitmap.close()
            ordered = reading_order(ordered, doc, page.width)
            ordered = merge_paragraph_lines(ordered, doc, page.width)
            section.children = [item[2] for item in ordered]
            add_page_context(section.children, doc)
            if len(lines if words else []) > 3 and any(
                abs(a[0]["top"] - b[0]["top"]) < 3 and b[0]["x0"] - a[-1]["x1"] >= 30
                for a, b in pairwise(lines)
            ):
                doc.warnings.append(
                    f"Trang {number + 1}: có nhiều cột; kiểm tra thứ tự đọc trong Review."
                )
            doc.blocks.append(section)
            page.close()
    doc.warnings = list(dict.fromkeys(doc.warnings))
    return doc


def reading_order(items: list, doc: Document, page_width: float) -> list:
    """Read ordinary two-column bands left-to-right, bounded by full-width blocks."""
    items = sorted(items, key=lambda item: (item[0], item[1]))
    lookup = doc.by_id()
    midpoint = page_width / 2

    def bounds(item):
        block = item[2]
        if block.meta.get("bbox"):
            return block.meta["bbox"]
        if block.segments:
            return lookup[block.segments[0]].anchor.get("bbox", [0, item[0], page_width, item[0]])
        return [0, item[0], page_width, item[0]]

    left = [item for item in items if bounds(item)[2] < midpoint - 10]
    right = [item for item in items if bounds(item)[0] > midpoint + 10]
    if len(left) < 2 or len(right) < 2:
        return items
    result, band = [], []

    def flush():
        result.extend(sorted(band, key=lambda item: (bounds(item)[0] > midpoint, item[0])))
        band.clear()

    for item in items:
        bbox = bounds(item)
        if bbox[0] <= midpoint <= bbox[2]:
            flush()
            result.append(item)
        else:
            band.append(item)
    flush()
    return result


def write_pdf_docx(doc: Document, target: Path):
    from docx import Document as WordDocument
    from docx.enum.section import WD_ORIENT, WD_SECTION_START
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    if not doc.complete():
        raise ValueError("Còn đoạn chưa dịch.")
    segments = doc.by_id()
    output = WordDocument()
    normal = output.styles["Normal"]
    normal.paragraph_format.space_after = Pt(0)
    normal.font.size = Pt(12)
    page_section = output.sections[0]
    page_meta = {}
    previous_bottom = None

    def add_text(paragraph, block):
        for sid in block.segments:
            segment = segments[sid]
            run = paragraph.add_run(segment.text)
            style = segment.style
            run.bold = bool(style.get("b"))
            run.italic = bool(style.get("i"))
            run.font.size = Pt(style.get("font_size", 12))
            font = style.get("font", "")
            families = {"Helvetica": "Arial", "Times": "Times New Roman", "Courier": "Courier New"}
            family = re.split(r"[-,]", font)[0]
            if family:
                run.font.name = families.get(family, family)
            run.font.color.rgb = RGBColor.from_string(style.get("color", "000000"))
        if not block.segments and block.meta.get("literal"):
            paragraph.add_run(str(block.meta["literal"]))

    def format_paragraph(paragraph, block, left_edge=None):
        bbox = block.meta.get("bbox")
        fmt = paragraph.paragraph_format
        fmt.space_after = Pt(0)
        fmt.widow_control = True
        if bbox:
            origin = page_meta.get("left_margin", 36) if left_edge is None else left_edge
            fmt.left_indent = Pt(max(0, bbox[0] - origin))
            gap = bbox[1] - previous_bottom if previous_bottom is not None else 0
            fmt.space_before = Pt(max(0, gap) * page_meta.get("gap_scale", 1))
        if block.meta.get("line_spacing"):
            size = max((segments[sid].style.get("font_size", 12) for sid in block.segments), default=12)
            fmt.line_spacing = Pt(max(block.meta["line_spacing"], size * 1.15))

    def render(block, container=None, left_edge=None, column_width=None):
        nonlocal previous_bottom
        container = output if container is None else container
        available = column_width or (
            page_section.page_width - page_section.left_margin - page_section.right_margin
        ) / 12700
        bbox = block.meta.get("bbox")
        if block.kind == "columns":
            # A borderless table keeps both columns side by side while letting
            # translated text reflow. A separate gutter preserves their origins.
            gap = bbox[1] - previous_bottom if previous_bottom is not None else 0
            if gap > 0:
                spacer = output.add_paragraph()
                spacer.paragraph_format.line_spacing = Pt(gap * page_meta.get("gap_scale", 1))
            left, right = [child.meta["bbox"] for child in block.children]
            origin = page_meta.get("left_margin", 36)
            gutter = min(36, right[0] - left[2])
            widths = [right[0] - origin - gutter, gutter, available - (right[0] - origin)]
            if min(widths[0], widths[2]) < 24:
                for column in block.children:
                    for child in column.children:
                        render(child)
                return
            table = output.add_table(rows=1, cols=3)
            table.autofit = False
            table.alignment = WD_TABLE_ALIGNMENT.LEFT
            for index, width in enumerate(widths):
                table.columns[index].width = Pt(width)
                cell = table.cell(0, index)
                cell.width = Pt(width)
                margins = OxmlElement("w:tcMar")
                for side in ("top", "left", "bottom", "right"):
                    edge = OxmlElement("w:" + side)
                    edge.set(qn("w:w"), "0")
                    edge.set(qn("w:type"), "dxa")
                    margins.append(edge)
                cell._tc.get_or_add_tcPr().append(margins)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            borders = OxmlElement("w:tblBorders")
            for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
                edge = OxmlElement("w:" + side)
                edge.set(qn("w:val"), "nil")
                borders.append(edge)
            table._tbl.tblPr.append(borders)
            for index, column in zip((0, 2), block.children, strict=True):
                cell = table.cell(0, index)
                empty = cell.paragraphs[0]._p
                cell._tc.remove(empty)
                previous_bottom = bbox[1]
                column_origin = origin if index == 0 else right[0]
                for child in column.children:
                    render(child, cell, column_origin, widths[index])
        elif block.kind == "table":
            width = max((len(row.children) for row in block.children), default=0)
            if width:
                gap = bbox[1] - previous_bottom if bbox and previous_bottom is not None else 0
                if gap > 0:
                    spacer = output.add_paragraph()
                    spacer.paragraph_format.line_spacing = Pt(gap * page_meta.get("gap_scale", 1))
                    spacer.paragraph_format.space_before = Pt(0)
                    spacer.paragraph_format.space_after = Pt(0)
                table = output.add_table(rows=len(block.children), cols=width)
                table.style = "Table Grid"
                table.autofit = False
                widths = []
                for column in range(width):
                    cell = next((row.children[column] for row in block.children if len(row.children) > column), None)
                    bounds = cell.meta.get("bbox") if cell else None
                    widths.append(bounds[2] - bounds[0] if bounds else available / width)
                scale = min(1, available / max(1, sum(widths)))
                for column, size in enumerate(widths):
                    table.columns[column].width = Pt(size * scale)
                for i, row in enumerate(block.children):
                    cell_boxes = [cell.meta["bbox"] for cell in row.children if cell.meta.get("bbox")]
                    if cell_boxes:
                        table.rows[i].height = Pt(max(bounds[3] - bounds[1] for bounds in cell_boxes))
                        table.rows[i].height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
                    for j, cell in enumerate(row.children):
                        target_cell = table.cell(i, j)
                        target_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                        target_cell.width = Pt(widths[j] * scale)
                        add_text(target_cell.paragraphs[0], cell)
        elif block.kind == "image":
            path = doc.assets.get(block.meta.get("asset"))
            if path:
                paragraph = container.add_paragraph()
                format_paragraph(paragraph, block, left_edge)
                indent = paragraph.paragraph_format.left_indent
                usable = max(1, available - ((indent or 0) / 12700))
                width = bbox[2] - bbox[0] if bbox else usable
                paragraph.add_run().add_picture(path, width=Pt(min(width, usable)))
        else:
            paragraph = container.add_paragraph(
                style=f"Heading {block.meta.get('level', 1)}" if block.kind == "heading" else None
            )
            add_text(paragraph, block)
            format_paragraph(paragraph, block, left_edge)
        if bbox:
            previous_bottom = bbox[3]

    pages = [block for block in doc.blocks if block.kind == "section"]
    for number, page in enumerate(pages):
        if number:
            page_section = output.add_section(WD_SECTION_START.NEW_PAGE)
            # The section-break paragraph is not source content. Its default
            # body-text height can otherwise create an extra nearly empty page.
            boundary = output.paragraphs[-1].paragraph_format
            boundary.line_spacing = Pt(1)
            boundary.space_before = Pt(0)
            boundary.space_after = Pt(0)
        page_meta = dict(page.meta)
        # Keep typography and object sizes; modestly reduce whitespace when
        # translating Chinese prose into longer English to limit orphan pages.
        grows = any(
            re.search(r"[\u3400-\u9fff]", s.source) and len(s.text) > len(s.source)
            for s in doc.segments if s.anchor.get("page") == page.meta.get("page")
        )
        page_meta["gap_scale"] = 0.9 if grows else 1
        if page_meta.get("width") and page_meta.get("height"):
            page_section.page_width = Pt(page_meta["width"])
            page_section.page_height = Pt(page_meta["height"])
            page_section.orientation = (
                WD_ORIENT.LANDSCAPE if page_meta["width"] > page_meta["height"]
                else WD_ORIENT.PORTRAIT
            )
        boxes = [child.meta["bbox"] for child in page.children if child.meta.get("bbox")]
        if boxes:
            page_meta["left_margin"] = min(72, max(18, min(bbox[0] for bbox in boxes)))
            right = page_meta.get("width", page_section.page_width / 12700) - max(bbox[2] for bbox in boxes)
            bottom = page_meta.get("height", page_section.page_height / 12700) - max(bbox[3] for bbox in boxes)
            page_section.left_margin = Pt(page_meta["left_margin"])
            page_section.right_margin = Pt(min(72, max(18, right)))
            page_section.top_margin = Pt(min(72, max(18, min(bbox[1] for bbox in boxes))))
            # English often needs more lines; use the source's remaining blank area
            # before spilling a figure/caption onto an otherwise empty extra page.
            page_section.bottom_margin = Pt(min(18, max(12, bottom)))
        previous_bottom = None
        for child in column_bands(page.children, doc, page_meta.get("width", 612)):
            render(child)
    if not pages:
        for block in doc.blocks:
            render(block)
    output.save(target)
