import io
import zipfile

import pytest
from lxml import etree

from converter.markdown import render_markdown
from converter.office import NS, read_office, write_office


def chart_workbook(path):
    import xlsxwriter
    workbook = xlsxwriter.Workbook(path)
    sheet = workbook.add_worksheet("测试数据")
    for cell, value in {"A1": "参数", "B1": "工作电压", "A2": "控制器", "A3": "传感器",
                        "B2": 12, "B3": 100, "C2": "=B2*2"}.items():
        sheet.write(cell, value)
    chart = workbook.add_chart({"type": "column"})
    chart.add_series({"name": "='测试数据'!$B$1", "categories": "='测试数据'!$A$2:$A$3",
                      "values": "='测试数据'!$B$2:$B$3"})
    chart.set_title({"name": "压力比较"})
    sheet.insert_chart("E2", chart)
    workbook.close()


def chart_without_text(data):
    root = etree.fromstring(data)
    for node in root.xpath(".//a:t | .//c:strCache/c:pt/c:v | .//c:strLit/c:pt/c:v", namespaces=NS):
        node.text = ""
    return etree.tostring(root, method="c14n")


def test_chart_labels_follow_reviewed_cells_while_styles_and_numeric_data_survive(tmp_path):
    from openpyxl import load_workbook
    source = tmp_path / "chart.xlsx"
    chart_workbook(source)
    doc = read_office(source, tmp_path / "assets")
    translations = {"参数": "Parameter", "工作电压": "Operating voltage", "控制器": "Controller",
                    "传感器": "Sensor", "压力比较": "Pressure comparison"}
    for segment in doc.segments:
        segment.target = translations.get(segment.source, segment.source)
    assert len([s for s in doc.segments if s.source == "控制器"]) == 1
    controller = next(s for s in doc.segments if s.source == "控制器")
    controller.target = "Reviewed controller"
    output = tmp_path / "translated.xlsx"
    write_office(doc, output)
    native = load_workbook(output)["测试数据"]
    assert native["A2"].value == "Reviewed controller"
    assert native["B2"].value == 12 and native["B3"].value == 100
    assert native["C2"].value == "=B2*2"
    with zipfile.ZipFile(source) as before, zipfile.ZipFile(output) as after:
        name = "xl/charts/chart1.xml"
        assert chart_without_text(before.read(name)) == chart_without_text(after.read(name))
        root = etree.fromstring(after.read(name))
        labels = root.xpath(".//c:strCache/c:pt/c:v/text()", namespaces=NS)
        assert labels == ["Operating voltage", "Reviewed controller", "Sensor"]
        assert root.xpath(".//c:title//a:t/text()", namespaces=NS) == ["Pressure comparison"]
        assert before.read("xl/styles.xml") == after.read("xl/styles.xml")
        assert all(before.read(name) == after.read(name) for name in before.namelist() if name.endswith(".rels"))
    assert "Pressure comparison" in render_markdown(doc, "assets")
    assert "压力比较" in render_markdown(doc, "assets", source=True)


def test_embedded_powerpoint_workbook_and_chart_cache_use_the_same_translation(tmp_path):
    from openpyxl import load_workbook
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches
    source = tmp_path / "chart.pptx"
    ppt = Presentation()
    slide = ppt.slides.add_slide(ppt.slide_layouts[6])
    data = CategoryChartData()
    data.categories = ["控制器", "传感器"]
    data.add_series("工作电压", (12, 100))
    chart = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1),
                                  Inches(8), Inches(5), data).chart
    chart.has_title = True
    chart.chart_title.text_frame.text = "压力比较"
    ppt.save(source)
    doc = read_office(source, tmp_path / "assets")
    translations = {"工作电压": "Operating voltage", "控制器": "Controller",
                    "传感器": "Sensor", "压力比较": "Pressure comparison"}
    for segment in doc.segments:
        segment.target = translations.get(segment.source, segment.source)
    assert len([s for s in doc.segments if s.source == "控制器"]) == 1
    controller = next(s for s in doc.segments if s.source == "控制器")
    controller.target = "Reviewed controller"
    output = tmp_path / "translated.pptx"
    write_office(doc, output)
    with zipfile.ZipFile(source) as before, zipfile.ZipFile(output) as after:
        name = "ppt/charts/chart1.xml"
        assert chart_without_text(before.read(name)) == chart_without_text(after.read(name))
        root = etree.fromstring(after.read(name))
        assert root.xpath(".//c:strCache/c:pt/c:v/text()", namespaces=NS) == [
            "Operating voltage", "Reviewed controller", "Sensor"]
        embedded = controller.anchor["embedded"]
        sheet = load_workbook(io.BytesIO(after.read(embedded))).active
        assert sheet["A2"].value == "Reviewed controller" and sheet["B1"].value == "Operating voltage"
        assert sheet["B2"].value == 12 and sheet["B3"].value == 100
        with zipfile.ZipFile(io.BytesIO(before.read(embedded))) as original, zipfile.ZipFile(io.BytesIO(after.read(embedded))) as translated:
            assert original.read("xl/styles.xml") == translated.read("xl/styles.xml")
            assert set(original.namelist()) == set(translated.namelist())
        assert before.read("ppt/slides/slide1.xml") == after.read("ppt/slides/slide1.xml")


def test_smartart_source_and_split_cache_share_a_target_without_duplicating_text(word_file, tmp_path):
    source = tmp_path / "diagram.docx"
    a = NS["a"]
    data = f'<dgm:dataModel xmlns:dgm="http://schemas.openxmlformats.org/drawingml/2006/diagram" xmlns:a="{a}"><dgm:ptLst><dgm:pt modelId="node-1"><dgm:t><a:p><a:r><a:t>控制器</a:t></a:r></a:p></dgm:t></dgm:pt></dgm:ptLst></dgm:dataModel>'
    drawing = f'<dsp:drawing xmlns:dsp="http://schemas.microsoft.com/office/drawing/2008/diagram" xmlns:a="{a}"><dsp:spTree><dsp:sp><dsp:txBody><a:p><a:r><a:t>控制</a:t></a:r><a:r><a:t>器</a:t></a:r></a:p></dsp:txBody></dsp:sp></dsp:spTree></dsp:drawing>'
    data = etree.tostring(etree.fromstring(data.encode()), pretty_print=True)
    with zipfile.ZipFile(word_file) as original, zipfile.ZipFile(source, "w") as new:
        for entry in original.infolist():
            new.writestr(entry, original.read(entry.filename))
        new.writestr("word/diagrams/data1.xml", data)
        new.writestr("word/diagrams/drawing1.xml", drawing)
    doc = read_office(source, tmp_path / "assets")
    assert len([s for s in doc.segments if s.source == "控制器"]) == 1
    for segment in doc.segments:
        segment.target = "Controller" if segment.source == "控制器" else segment.source
    output = tmp_path / "translated.docx"
    write_office(doc, output)
    with zipfile.ZipFile(output) as archive:
        for part in ["word/diagrams/data1.xml", "word/diagrams/drawing1.xml"]:
            root = etree.fromstring(archive.read(part))
            assert "".join(root.xpath(".//a:t/text()", namespaces=NS)) == "Controller"


@pytest.mark.parametrize("older_version", [1, 2, 3])
def test_changed_office_adapter_rejects_resume_but_preserves_historical_exports(word_file, store, tmp_path, older_version):
    from converter.pipeline import Pipeline
    job = store.add(word_file)
    pipeline = Pipeline(store)
    pipeline.run(job, None, "markdown", tmp_path / "out", operation="convert")
    doc = store.load_document(job)
    doc.adapter_version = older_version
    store.save_document(job, doc)
    with pytest.raises(ValueError, match="adapter"):
        pipeline.parse(job, "eng")
    result = pipeline.run(job, None, "markdown", tmp_path / "out", operation="convert", export_only=True)
    assert result["markdown"]["status"] == "done"


def test_smartart_inline_breaks_use_retained_segment_ids_after_cache_binding(tmp_path):
    from docx import Document as WordDocument

    from converter.markdown import render_markdown

    original, source = tmp_path / "original.docx", tmp_path / "diagram.docx"
    WordDocument().save(original)
    diagram = (f'<root xmlns:a="{NS["a"]}"><a:p><a:r><a:t>控制器</a:t></a:r>'
               '<a:br/><a:r><a:t>传感器</a:t></a:r></a:p></root>')
    with zipfile.ZipFile(original) as before, zipfile.ZipFile(source, "w") as after:
        for entry in before.infolist():
            after.writestr(entry, before.read(entry.filename))
        for part in ["word/diagrams/data1.xml", "word/diagrams/drawing1.xml"]:
            after.writestr(part, diagram)
    doc = read_office(source, tmp_path / "assets")
    assert len(doc.segments) == 2
    assert "控制器<br>传感器" in render_markdown(doc, "assets", source=True)
    for segment in doc.segments:
        segment.target = {"控制器": "Controller", "传感器": "Sensor"}[segment.source]
    assert "Controller<br>Sensor" in render_markdown(doc, "assets")


def test_external_chart_data_is_not_bound_to_a_same_named_local_sheet(tmp_path):
    original = tmp_path / "original.xlsx"
    chart_workbook(original)
    source = tmp_path / "external.xlsx"
    with zipfile.ZipFile(original) as before, zipfile.ZipFile(source, "w") as after:
        for entry in before.infolist():
            data = before.read(entry.filename)
            if entry.filename == "xl/charts/chart1.xml":
                root = etree.fromstring(data)
                node = etree.SubElement(root, "{" + NS["c"] + "}externalData")
                node.set("{" + NS["r"] + "}id", "rIdExternal")
                data = etree.tostring(root)
            after.writestr(entry, data)
        after.writestr("xl/charts/_rels/chart1.xml.rels", f'<Relationships xmlns="{NS["rel"]}"><Relationship Id="rIdExternal" Target="https://example.invalid/data.xlsx" TargetMode="External" Type="{NS["r"]}/package"/></Relationships>')
    doc = read_office(source, tmp_path / "assets")
    controllers = [s for s in doc.segments if s.source == "控制器"]
    assert len(controllers) == 2
    for segment in doc.segments:
        segment.target = segment.source
    next(s for s in controllers if "cell" in s.anchor).target = "Local reviewed controller"
    next(s for s in controllers if "paths" in s.anchor).target = "External cached controller"
    output = tmp_path / "translated.xlsx"
    write_office(doc, output)
    with zipfile.ZipFile(output) as archive:
        root = etree.fromstring(archive.read("xl/charts/chart1.xml"))
        assert "External cached controller" in root.xpath(".//c:strCache/c:pt/c:v/text()", namespaces=NS)
        assert "Local reviewed controller" not in root.xpath(".//c:strCache/c:pt/c:v/text()", namespaces=NS)
    assert any("refresh" in warning for warning in doc.warnings)


def test_chart_part_discovery_uses_content_type_for_custom_file_names(tmp_path):
    original = tmp_path / "original.xlsx"
    chart_workbook(original)
    source = tmp_path / "custom.xlsx"
    with zipfile.ZipFile(original) as before, zipfile.ZipFile(source, "w") as after:
        for entry in before.infolist():
            name = entry.filename
            data = before.read(name)
            if name == "xl/charts/chart1.xml":
                name = "xl/charts/automotive-plot.xml"
            elif name == "[Content_Types].xml" or name.endswith(".rels"):
                data = data.replace(b"chart1.xml", b"automotive-plot.xml")
            after.writestr(name, data)
    doc = read_office(source, tmp_path / "assets")
    title = next((s for s in doc.segments if s.source == "压力比较"), None)
    assert title is not None
    assert title.anchor["part"] == "xl/charts/automotive-plot.xml"
    for segment in doc.segments:
        segment.target = "Pressure comparison" if segment is title else segment.source
    output = tmp_path / "translated.xlsx"
    write_office(doc, output)
    with zipfile.ZipFile(output) as archive:
        root = etree.fromstring(archive.read("xl/charts/automotive-plot.xml"))
        assert root.xpath(".//a:t/text()", namespaces=NS) == ["Pressure comparison"]
