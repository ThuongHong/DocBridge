import zipfile

import openpyxl
import pytest
from lxml import etree

from converter.markdown import render_markdown
from converter.office import NS, read_office, write_office


def test_sheet_titles_headers_notes_and_references(tmp_path):
    source = tmp_path / 'source.xlsx'
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = '中文表'
    sheet['A1'] = '控制器'
    sheet['A1'].comment = openpyxl.comments.Comment('检查连接', '作者')
    sheet.oddHeader.center.text = '报告 &P / &N'
    sheet['B1'] = '=\'中文表\'!A1&"中文表!A1"'
    sheet['C1'].hyperlink = '#\'中文表\'!A1'
    book.save(source)
    doc = read_office(source, tmp_path / 'assets')
    sources = [s.source for s in doc.segments]
    assert '中文表' in sources
    assert '检查连接' in sources
    assert any('报告 ' == s for s in sources)
    for segment in doc.segments:
        segment.target = {'中文表': 'Report', '控制器': 'Controller', '检查连接': 'Check connection',
                          '报告 ': 'Report '}.get(segment.source, segment.source)
    output = tmp_path / 'out.xlsx'
    write_office(doc, output)
    result = openpyxl.load_workbook(output)
    assert result.sheetnames == ['Report']
    assert result.active['B1'].value == '=\'Report\'!A1&"中文表!A1"'
    assert result.active['A1'].comment.text == 'Check connection'
    assert result.active.oddHeader.center.text == 'Report &P / &N'
    assert result.active['C1'].hyperlink.target == '#\'Report\'!A1'
    assert '## Report' in render_markdown(doc, 'assets')
    assert '## 中文表' in render_markdown(doc, 'assets', source=True)


@pytest.mark.parametrize('shared', [False, True])
def test_rich_groups_and_phonetics(tmp_path, shared):
    source = tmp_path / 'rich.xlsx'
    book = openpyxl.Workbook()
    book.active['A1'] = 'placeholder'
    book.save(source)
    patched = tmp_path / 'patched.xlsx'
    content = '<r><rPr><b/></rPr><t>控</t></r><r><rPr><b/></rPr><t>制</t></r><r><rPr><b/></rPr><t>器</t></r><r><rPr><color rgb="FFFF0000"/></rPr><t>检查</t></r><rPh sb="0" eb="3"><t>kong zhi qi</t></rPh>'
    with zipfile.ZipFile(source) as before, zipfile.ZipFile(patched, 'w') as after:
        for entry in before.infolist():
            data = before.read(entry.filename)
            if entry.filename == 'xl/worksheets/sheet1.xml':
                root = etree.fromstring(data)
                cell = root.find('.//s:c', NS)
                for child in list(cell):
                    cell.remove(child)
                cell.set('t', 's' if shared else 'inlineStr')
                if shared:
                    etree.SubElement(cell, '{'+NS['s']+'}v').text = '0'
                else:
                    cell.append(etree.fromstring(f'<is xmlns="{NS["s"]}">{content}</is>'))
                data = etree.tostring(root)
            after.writestr(entry, data)
        if shared:
            after.writestr('xl/sharedStrings.xml', f'<sst xmlns="{NS["s"]}" count="1" uniqueCount="1"><si>{content}</si></sst>')
    doc = read_office(patched, tmp_path / 'assets')
    cells = [s for s in doc.segments if 'cell' in s.anchor]
    assert [s.source for s in cells] == ['控制器', '检查']
    assert all('kong' not in s.context for s in cells)
    for segment in doc.segments:
        segment.target = {'控制器': 'Controller', '检查': 'Check'}.get(segment.source, segment.source)
    output = tmp_path / 'out.xlsx'
    write_office(doc, output)
    reread = read_office(output, tmp_path / 'assets2')
    assert [s.source for s in reread.segments if 'cell' in s.anchor] == ['Controller', 'Check']


def test_sheet_reference_lexer_preserves_strings_external_and_structured_labels():
    from converter.excel_text import rewrite_references

    expression = "SUM('中文表:尾表'!A1,中文表!A1,[Book.xlsx]中文表!A1,'[Book.xlsx]中文表'!A1,Table[中文表!A1],\"中文表!A1\")"
    assert rewrite_references(expression, {'中文表': "Bob's report", '尾表': 'End'}) == (
        "SUM('Bob''s report:End'!A1,'Bob''s report'!A1,[Book.xlsx]中文表!A1,'[Book.xlsx]中文表'!A1,Table[中文表!A1],\"中文表!A1\")")
    with pytest.raises(ValueError, match='INDIRECT.*Preserve'):
        rewrite_references('INDIRECT("中文"&"表!A1")', {'中文表': 'Report'})
    assert rewrite_references('INDIRECT("中文表!A1")', {'中文表': '中文表'}) == 'INDIRECT("中文表!A1")'


def test_sheet_name_sanitization_reserves_existing_names():
    from converter.excel_text import sheet_names
    from converter.models import Segment

    segments = [Segment('1', '中文', {'category': 'sheet_title'}, target="'A/B:*?[]\\'"),
                Segment('2', '另一表', {'category': 'sheet_title'}, target='a b     '),
                Segment('3', '长表', {'category': 'sheet_title'}, target='x' * 40),
                Segment('4', '长表二', {'category': 'sheet_title'}, target='X' * 40)]
    names = sheet_names(segments, ['A B'])
    assert len({name.casefold() for name in names.values()}) == len(names)
    assert all(0 < len(name) <= 31 and not any(c in name for c in '/*?:[]\\') for name in names.values())


def test_header_controls_with_chinese_font_names_are_not_segments(tmp_path):
    source = tmp_path / 'header.xlsx'
    book = openpyxl.Workbook()
    book.active.oddHeader.center.text = '&"宋体,Bold"&14&KFF0000报告&P+1&N&D&T&F&A&&'
    book.save(source)
    doc = read_office(source, tmp_path / 'assets')
    assert [s.source for s in doc.segments] == ['报告']
    doc.segments[0].target = 'Report & notes'
    output = tmp_path / 'out.xlsx'
    write_office(doc, output)
    with zipfile.ZipFile(output) as archive:
        text = etree.fromstring(archive.read('xl/worksheets/sheet1.xml')).find('s:headerFooter/s:oddHeader', NS).text
        assert '&"宋体,Bold"' in text
        assert 'Report && notes&P+1&N&D&T&F&A&&' in text


def test_export_coverage_ignores_orphan_shared_strings_and_reports_formulas(spreadsheet, tmp_path):
    doc = read_office(spreadsheet, tmp_path / 'assets')
    for segment in doc.segments:
        segment.target = 'English'
    output = tmp_path / 'out.xlsx'
    write_office(doc, output)
    assert 'translatable_cell' not in doc.coverage['han_locations']
    assert doc.coverage['images_not_ocr_scanned'] == 1
    assert all(len(v['examples']) <= 5 for v in doc.coverage['han_locations'].values())


def test_pipeline_reuse_keeps_contexts_and_reviewed_targets_separate(tmp_path, store):
    from converter.pipeline import Pipeline

    source = tmp_path / 'repeat.xlsx'
    book = openpyxl.Workbook()
    sheet = book.active
    for index in range(1, 12):
        sheet.cell(index, 1, '控制器')
    book.create_sheet('Other')['A1'] = '控制器'
    book.save(source)
    job = store.add(source)
    pipeline = Pipeline(store)
    doc = pipeline.parse(job, 'eng')
    doc.segments[0].target = 'Reviewed override'
    store.save_document(job, doc)

    class Provider:
        batch_chars = 1000
        calls = []

        def fingerprint(self, glossary):
            return 'profile'

        def translate(self, segments, glossary):
            self.calls.extend(segments)
            return {s.id: 'Controller' for s in segments}

    provider = Provider()
    result = pipeline.translate(job, provider, {}, 'eng')
    assert len(provider.calls) == 2
    assert result.segments[0].target == 'Reviewed override'
    assert all(s.target == 'Controller' for s in result.segments[1:])
    result.segments[1].target = 'Independent review'
    output = tmp_path / 'out.xlsx'
    write_office(result, output)
    sheet = openpyxl.load_workbook(output).active
    assert sheet['A1'].value == 'Reviewed override'
    assert sheet['A2'].value == 'Independent review'
    assert sheet['A3'].value == 'Controller'


@pytest.mark.parametrize('expression,expected', [
    ("SUM('数据一':'数据二'!A1)", "SUM('One:Two'!A1)"),
    ("SUM('数据一':数据二!A1)", "SUM('One:Two'!A1)"),
    ("SUM(数据一:'数据二'!A1)", "SUM('One:Two'!A1)"),
    ('''SUM('数据"引号'!A1,"数据一!A1")''', '''SUM('Quoted'!A1,"数据一!A1")'''),
])
def test_quoted_3d_endpoints_and_double_quotes_in_sheet_names(expression, expected):
    from converter.excel_text import rewrite_references

    assert rewrite_references(expression, {'数据一': 'One', '数据二': 'Two', '数据"引号': 'Quoted'}) == expected


def test_translation_reuse_admits_frequent_strings_after_cache_capacity(spreadsheet, tmp_path, store):
    from converter.models import Segment
    from converter.pipeline import Pipeline

    job = store.add(spreadsheet)
    pipeline = Pipeline(store)
    doc = pipeline.parse(job, 'eng')
    doc.segments = [Segment(str(i), f'不同文本{i}', context='same') for i in range(4200)]
    doc.segments.extend(Segment(f'repeat{i}', '0-无', context='same') for i in range(100))
    store.save_document(job, doc)

    class Provider:
        batch_chars = 10_000_000
        handles_long_segments = True
        count = 0

        def fingerprint(self, glossary):
            return 'profile'

        def translate(self, segments, glossary):
            self.count += len(segments)
            return {s.id: 'English' for s in segments}

    provider = Provider()
    assert pipeline.translate(job, provider, {}, 'eng').complete()
    assert provider.count == 4201


def test_hyperlink_constant_targets_follow_sheet_renames_and_dynamic_targets_fail():
    from converter.excel_text import rewrite_references

    names = {'中文表': 'Report'}
    assert rewrite_references('HYPERLINK("#\'中文表\'!A1","中文表!A1")', names) == (
        'HYPERLINK("#\'Report\'!A1","中文表!A1")')
    assert rewrite_references('HYPERLINK("https://example.test/中文表!A1","link")', names) == (
        'HYPERLINK("https://example.test/中文表!A1","link")')
    for expression in ['HYPERLINK(A1,"link")', 'HYPERLINK("#中文表!"&A1,"link")']:
        with pytest.raises(ValueError, match='HYPERLINK.*Preserve'):
            rewrite_references(expression, names)


def test_pivot_cache_local_source_sheet_and_xml_comments_survive_renaming(tmp_path):
    from converter.excel_text import rename_package
    from converter.office import rels, xml

    source = tmp_path / 'pivot.xlsx'
    book = openpyxl.Workbook()
    book.active.title = '中文表'
    book.save(source)
    with zipfile.ZipFile(source, 'a') as archive:
        archive.writestr('xl/pivotCache/pivotCacheDefinition1.xml',
                         f'<pivotCacheDefinition xmlns="{NS["s"]}"><!--valid--><cacheSource type="worksheet"><worksheetSource sheet="中文表" ref="A1:B5"/></cacheSource></pivotCacheDefinition>')
    with zipfile.ZipFile(source) as archive:
        roots = {}
        rename_package(archive, roots, {'中文表': 'Report'}, 'xl/workbook.xml', rels, xml)
    root = roots['xl/pivotCache/pivotCacheDefinition1.xml']
    assert root.find('s:cacheSource/s:worksheetSource', NS).get('sheet') == 'Report'
    assert isinstance(root[0], etree._Comment)


def test_embedded_chart_sheet_titles_and_outer_formula_refs_stay_consistent(tmp_path):
    import io

    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches

    original, source = tmp_path / 'original.pptx', tmp_path / 'source.pptx'
    ppt = Presentation()
    slide = ppt.slides.add_slide(ppt.slide_layouts[6])
    data = CategoryChartData()
    data.categories = ['控制器']
    data.add_series('数值', [1])
    slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1), Inches(5), Inches(4), data)
    ppt.save(original)
    with zipfile.ZipFile(original) as before, zipfile.ZipFile(source, 'w') as after:
        for entry in before.infolist():
            content = before.read(entry.filename)
            if entry.filename.endswith('.xlsx'):
                buffer = io.BytesIO()
                with zipfile.ZipFile(io.BytesIO(content)) as book, zipfile.ZipFile(buffer, 'w') as changed:
                    for item in book.infolist():
                        payload = book.read(item.filename)
                        if item.filename == 'xl/workbook.xml':
                            root = etree.fromstring(payload)
                            root.find('s:sheets/s:sheet', NS).set('name', '中文表')
                            payload = etree.tostring(root)
                        changed.writestr(item, payload)
                content = buffer.getvalue()
            elif entry.filename == 'ppt/charts/chart1.xml':
                root = etree.fromstring(content)
                for formula in root.findall('.//c:f', NS):
                    formula.text = formula.text.replace('Sheet1', '中文表')
                content = etree.tostring(root)
            after.writestr(entry, content)
    doc = read_office(source, tmp_path / 'assets')
    assert '中文表' in [s.source for s in doc.segments]
    for segment in doc.segments:
        segment.target = {'中文表': 'Report', '控制器': 'Controller', '数值': 'Value'}.get(segment.source, segment.source)
    output = tmp_path / 'out.pptx'
    write_office(doc, output)
    with zipfile.ZipFile(output) as archive:
        root = etree.fromstring(archive.read('ppt/charts/chart1.xml'))
        assert all(f.text.startswith("'Report'!") for f in root.findall('.//c:f', NS))
        embedded = next(name for name in archive.namelist() if name.endswith('.xlsx'))
        assert openpyxl.load_workbook(io.BytesIO(archive.read(embedded))).sheetnames == ['Report']
    assert '## Chart data: Report' in render_markdown(doc, 'assets')
    assert '## Chart data: 中文表' in render_markdown(doc, 'assets', source=True)


def test_coverage_distinguishes_formulas_formats_and_unsupported_prompts(tmp_path):
    from openpyxl.worksheet.datavalidation import DataValidation

    source = tmp_path / 'coverage.xlsx'
    book = openpyxl.Workbook()
    sheet = book.active
    sheet['A1'] = '控制器'
    sheet['B1'] = '=IF(A1="控制器","正常","故障")'
    sheet['C1'] = 1
    sheet['C1'].number_format = '0"元"'
    validation = DataValidation(type='whole', operator='between', formula1=1, formula2=10)
    validation.prompt = '输入数量'
    sheet.add_data_validation(validation)
    validation.add(sheet['C1'])
    book.save(source)
    doc = read_office(source, tmp_path / 'assets')
    for segment in doc.segments:
        segment.target = 'Controller'
    output = tmp_path / 'out.xlsx'
    write_office(doc, output)
    locations = doc.coverage['han_locations']
    assert 'translatable_cell' not in locations
    assert locations['preserved_formula']['count'] == 1
    assert locations['preserved_format_code']['count'] == 1
    assert locations['unsupported_user_visible_attribute']['count'] == 1


def dropdown_fixture(tmp_path):
    from openpyxl.worksheet.datavalidation import DataValidation

    source = tmp_path / 'dropdown.xlsx'
    book = openpyxl.Workbook()
    sheet = book.active
    sheet['A1'] = '0-无'
    sheet['A2'] = '1-内部'
    sheet['B1'] = '0-无'
    validation = DataValidation(type='list', formula1='"0-无,1-内部,2-秘密"')
    sheet.add_data_validation(validation)
    validation.add('A1:A20')
    ranged = DataValidation(type='list', formula1='$B$1:$B$10')
    sheet.add_data_validation(ranged)
    ranged.add('C1:C20')
    book.save(source)
    return source


def test_literal_dropdown_options_translate_and_reference_validations_stay_formulas(tmp_path):
    source = dropdown_fixture(tmp_path)
    doc = read_office(source, tmp_path / 'assets')
    options = [s for s in doc.segments if s.anchor.get('category') == 'dropdown_option']
    assert [s.source for s in options] == ['0-无', '1-内部', '2-秘密']
    translations = {'0-无': '0-None', '1-内部': '1-Internal', '2-秘密': '2-Secret'}
    for segment in doc.segments:
        segment.target = translations.get(segment.source, segment.source)
    output = tmp_path / 'out.xlsx'
    write_office(doc, output)
    sheet = openpyxl.load_workbook(output).active
    assert sheet['A1'].value == '0-None'
    assert [v.formula1 for v in sheet.data_validations.dataValidation] == ['"0-None,1-Internal,2-Secret"', '$B$1:$B$10']
    assert 'translatable_dropdown_option' not in doc.coverage['han_locations']
    assert '0-None' in render_markdown(doc, 'assets')
    assert '0-无' in render_markdown(doc, 'assets', source=True)


@pytest.mark.parametrize('invalid,message', [('Comma,option', 'commas'), ('x' * 255, '255'), ('0-Different', 'different translations')])
def test_invalid_or_unsynchronized_dropdown_targets_require_review(tmp_path, invalid, message):
    source = dropdown_fixture(tmp_path)
    doc = read_office(source, tmp_path / 'assets')
    for segment in doc.segments:
        segment.target = {'0-无': '0-None', '1-内部': '1-Internal', '2-秘密': '2-Secret'}.get(segment.source, segment.source)
    option = next(s for s in doc.segments if s.anchor.get('category') == 'dropdown_option' and s.source == '0-无')
    option.target = invalid
    with pytest.raises(ValueError, match=message):
        write_office(doc, tmp_path / 'out.xlsx')


def test_preserved_dropdown_options_still_check_changed_validated_cells(tmp_path):
    source = dropdown_fixture(tmp_path)
    doc = read_office(source, tmp_path / 'assets')
    for segment in doc.segments:
        segment.target = '0-None' if segment.anchor.get('cell') == 'A1' else segment.source
    with pytest.raises(ValueError, match='different translations'):
        write_office(doc, tmp_path / 'out.xlsx')


def test_reused_cell_targets_are_checkpointed_before_pause_and_resume(tmp_path, store):
    from converter.pipeline import Paused, Pipeline

    source = tmp_path / 'pause.xlsx'
    book = openpyxl.Workbook()
    for row in range(1, 21):
        book.active.cell(row, 1, '控制器')
    book.active['B20'] = '传感器'
    book.save(source)
    job = store.add(source)
    state = {'pause': False}

    class Provider:
        batch_chars = 60
        handles_long_segments = True
        calls = 0

        def fingerprint(self, glossary):
            return 'same-profile'

        def translate(self, segments, glossary):
            self.calls += 1
            return {s.id: 'English' for s in segments}

    def progress(job_id, message, done, total):
        if message == 'Đang dịch' and done:
            state['pause'] = True

    def control():
        if state['pause']:
            raise Paused()

    provider = Provider()
    with pytest.raises(Paused):
        Pipeline(store, progress=progress, control=control).translate(job, provider, {}, 'eng')
    saved = store.load_document(job)
    assert sum(s.target is not None for s in saved.segments) == 20
    assert not saved.complete()
    assert provider.calls == 1
    result = Pipeline(store).translate(job, provider, {}, 'eng')
    assert result.complete()
    assert provider.calls == 2
