import zipfile

import pytest
from lxml import etree

from converter.markdown import render_markdown
from converter.office import NS, read_office, write_office


@pytest.mark.parametrize("format", ["docx", "pptx"])
def test_breaks_and_tabs_survive_source_markdown_context_and_office_roundtrip(tmp_path, format):
    source = tmp_path / ("controls." + format)
    if format == "docx":
        from docx import Document
        from docx.shared import Inches
        book = Document()
        paragraph = book.add_paragraph()
        paragraph.paragraph_format.tab_stops.add_tab_stop(Inches(1))
        run = paragraph.add_run("控制器")
        run.bold = True
        run.add_tab()
        run.add_text("传感器")
        run.add_break()
        run.add_text("检查连接。")
        part, prefix = "word/document.xml", "w"
        expected_context = "控制器\t传感器\n检查连接。"
        expected_markdown = "**控制器**\t**传感器**<br>**检查连接。**"
    else:
        from pptx import Presentation
        from pptx.util import Inches
        book = Presentation()
        slide = book.slides.add_slide(book.slide_layouts[6])
        frame = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(3)).text_frame
        paragraph = frame.paragraphs[0]
        run = paragraph.add_run()
        run.text, run.font.bold = "控制器", True
        paragraph.add_line_break()
        run = paragraph.add_run()
        run.text, run.font.italic = "检查连接。", True
        part, prefix = "ppt/slides/slide1.xml", "a"
        expected_context = "控制器\n检查连接。"
        expected_markdown = "**控制器**<br>*检查连接。*"
    book.save(source)
    doc = read_office(source, tmp_path / "assets")
    assert {s.context for s in doc.segments} == {expected_context}
    assert expected_markdown in render_markdown(doc, "assets", source=True)
    for segment in doc.segments:
        segment.target = {"控制器": "Controller", "传感器": "Sensor",
                          "检查连接。": "Check the connection."}[segment.source]
    output = tmp_path / ("output." + format)
    write_office(doc, output)
    # Translation is the only allowed XML change: tab stops, actual tabs,
    # explicit line breaks, runs, properties and shape geometry must remain.
    with zipfile.ZipFile(source) as before, zipfile.ZipFile(output) as after:
        assert before.namelist() == after.namelist()
        for name in before.namelist():
            if name == part:
                left, right = etree.fromstring(before.read(name)), etree.fromstring(after.read(name))
                for tree in (left, right):
                    for node in tree.findall(f".//{prefix}:t", NS):
                        node.text = ""
                assert etree.tostring(left, method="c14n") == etree.tostring(right, method="c14n")
            else:
                assert before.read(name) == after.read(name)
    assert expected_markdown in render_markdown(doc, "assets", source=True)
    assert "<br>" in render_markdown(doc, "assets")
