import pytest
from openpyxl import load_workbook

from converter.excel_formulas import align_formula
from converter.office import read_office, write_office


def test_criteria_literals_follow_translated_labels_and_keep_numeric_formula_structure(tmp_path):
    import xlsxwriter
    source = tmp_path / "criteria.xlsx"
    book = xlsxwriter.Workbook(source)
    sheet = book.add_worksheet("数据")
    sheet.write("A1", "控制器")
    sheet.write("A2", "传感器")
    sheet.write("B1", 12)
    sheet.write("B2", 100)
    sheet.write_formula("C1", '=SUMIF(A1:A2,"控制器",B1:B2)', None, 12)
    sheet.write_formula("C2", '=COUNTIF(A1:A2,"传感器")', None, 1)
    sheet.write_formula("D1", "=B1*2", None, 24)
    book.close()
    doc = read_office(source, tmp_path / "assets")
    for segment in doc.segments:
        segment.target = {"控制器": "Controller", "传感器": "Sensor"}.get(segment.source, segment.source)
    output = tmp_path / "result.xlsx"
    write_office(doc, output)
    sheet = load_workbook(output)["数据"]
    assert sheet["C1"].value == '=SUMIF(A1:A2,"Controller",B1:B2)'
    assert sheet["C2"].value == '=COUNTIF(A1:A2,"Sensor")'
    assert sheet["D1"].value == "=B1*2"
    assert sheet["B1"].value == 12 and sheet["B2"].value == 100


def test_nested_criteria_preserve_whitespace_quotes_and_sheet_scope():
    cells = {"数据's": [(1, 1, '控制器"A', 'Controller"A'),
                          (1, 2, "传感器", "Sensor")],
             "other": [(1, 1, '控制器"A', "Different")]}
    formula = ('SUM(  SUMIF(\'数据\'\'s\'!$A:$A, "控制器""A",B:B),'
               'COUNTIFS(\'数据\'\'s\'!A1:A2,"<>传感器",B1:B2,">0"))')
    assert align_formula(formula, "other", cells) == (
        'SUM(  SUMIF(\'数据\'\'s\'!$A:$A, "Controller""A",B:B),'
        'COUNTIFS(\'数据\'\'s\'!A1:A2,"<>Sensor",B1:B2,">0"))')
    for formula in ['HYPERLINK("控制器.xlsx","控制器")', 'INDIRECT("控制器!A1")',
                    'COUNTIF(A:A,"控制器*")', 'COUNTIF(named_range,"控制器")',
                    'COUNTIF(A:A,">控制器")', 'B1*2']:
        assert align_formula(formula, "other", cells) == formula


@pytest.mark.parametrize("records", [
    [(1, 1, "控制器", "Controller"), (1, 2, "控制器", "Control unit")],
    [(1, 1, "控制器", "Controller"), (1, 2, "传感器", "Controller")],
    [(1, 1, "控制器", "Controller*")],
])
def test_ambiguous_or_pattern_creating_reviews_refuse_unsafe_criteria(records):
    with pytest.raises(ValueError):
        align_formula('COUNTIF(A:A,"控制器")', "数据", {"数据": records})


def test_existing_english_criterion_does_not_gain_new_matches_after_translation():
    with pytest.raises(ValueError):
        align_formula('COUNTIF(A:A,"Sensor")', "数据",
                      {"数据": [(1, 1, "传感器", "Sensor")]})


def test_reviewed_criteria_are_reexported_from_original_and_fail_without_partial_file(tmp_path):
    from openpyxl import Workbook
    source = tmp_path / "source.xlsx"
    book = Workbook()
    sheet = book.active
    sheet["A1"] = sheet["A2"] = "控制器"
    sheet["B1"] = '=COUNTIF(A1:A2,"控制器")'
    book.save(source)
    doc = read_office(source, tmp_path / "assets")
    for segment in doc.segments:
        segment.target = "Controller"
    output = tmp_path / "output.xlsx"
    write_office(doc, output)
    assert load_workbook(output).active["B1"].value == '=COUNTIF(A1:A2,"Controller")'
    for segment in doc.segments:
        segment.target = "Control unit"
    write_office(doc, output)
    assert load_workbook(output).active["B1"].value == '=COUNTIF(A1:A2,"Control unit")'
    doc.segments[0].target = "Different"
    untouched = tmp_path / "unsafe.xlsx"
    with pytest.raises(ValueError):
        write_office(doc, untouched)
    assert not untouched.exists()
