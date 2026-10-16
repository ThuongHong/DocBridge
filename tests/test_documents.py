import zipfile
from pathlib import Path

import pytest
from lxml import etree

from converter.markdown import render_markdown
from converter.models import file_hash
from converter.office import NS, read_office, write_office
from converter.pdf import read_pdf, write_pdf_docx
from converter.pipeline import Pipeline


def assert_unmodified_parts(source, target, allowed):
    with zipfile.ZipFile(source) as original, zipfile.ZipFile(target) as translated:
        assert set(original.namelist()) == set(translated.namelist())
        for part in original.namelist():
            if part not in allowed:
                assert original.read(part) == translated.read(part), part


def test_word_preserves_run_styles_headers_media_and_tables(word_file, tmp_path):
    from docx import Document

    doc = read_office(word_file, tmp_path / "assets")
    assert "测试文本" in [s.source for s in doc.segments]
    assert {"页眉", "页脚", "重要 ", "普通"} <= {s.source for s in doc.segments}
    for segment in doc.segments:
        segment.target = "English " + segment.source
    output = tmp_path / "translated.docx"
    write_office(doc, output)
    assert_unmodified_parts(word_file, output, {s.anchor["part"] for s in doc.segments})
    result = Document(output)
    assert result.paragraphs[1].runs[0].bold is True
    assert result.paragraphs[1].runs[1].italic is True
    assert result.sections[0].header.paragraphs[0].text == "English 页眉"
    assert len(result.tables) == 1
    assert len(result.inline_shapes) == 1
    markdown = render_markdown(doc, "translated.assets")
    assert "**English 重要**" in markdown
    assert "*English 普通*" in markdown
    assert "translated.assets/image-" in markdown


def test_powerpoint_group_table_notes_and_styles(slides_file, tmp_path):
    from pptx import Presentation

    doc = read_office(slides_file, tmp_path / "assets")
    assert {"粗体", "文本", "列名", "数据", "演讲备注"} <= {s.source for s in doc.segments}
    for segment in doc.segments:
        segment.target = "English " + segment.source
    output = tmp_path / "translated.pptx"
    write_office(doc, output)
    assert_unmodified_parts(slides_file, output, {s.anchor["part"] for s in doc.segments})
    result = Presentation(output)
    runs = result.slides[0].shapes[0].shapes[0].text_frame.paragraphs[0].runs
    assert runs[0].font.bold is True
    assert runs[1].font.italic is True
    assert "English 演讲备注" in result.slides[0].notes_slide.notes_text_frame.text


def test_excel_shared_strings_are_cell_specific_and_objects_survive(spreadsheet, tmp_path):
    from openpyxl import load_workbook

    doc = read_office(spreadsheet, tmp_path / "assets")
    same = [s for s in doc.segments if s.source == "相同"]
    assert len(same) == 2
    for segment in doc.segments:
        segment.target = "English " + segment.source
    same[0].target = "First context"
    same[1].target = "Second context"
    output = tmp_path / "translated.xlsx"
    write_office(doc, output)
    assert_unmodified_parts(
        spreadsheet, output, {"xl/worksheets/sheet1.xml", "xl/sharedStrings.xml",
                              "xl/workbook.xml", "xl/charts/chart1.xml"}
    )
    sheet = load_workbook(output)["English 原始名称"]
    assert sheet["A1"].value == "First context"
    assert sheet["A2"].value == "Second context"
    assert sheet["B1"].value == 123
    assert sheet["B2"].value == "=B1*2"
    assert len(sheet._charts) == 1
    assert len(sheet._images) == 1
    with zipfile.ZipFile(output) as archive:
        root = etree.fromstring(archive.read("xl/sharedStrings.xml"))
        rich = [s for s in root if s.find("s:r/s:rPr/s:b", NS) is not None]
        assert rich
        assert (
            rich[-1].find("s:r/s:rPr/s:i", NS) is not None
            or rich[-1].findall("s:r", NS)[1].find("s:rPr/s:i", NS) is not None
        )


@pytest.mark.parametrize(
    "mode,extensions", [("office", {".docx"}), ("markdown", {".md"}), ("both", {".docx", ".md"})]
)
def test_modes_and_review_use_one_translation(
    word_file, store, provider, tmp_path, monkeypatch, mode, extensions
):
    job = store.add(word_file)
    pipeline = Pipeline(store)
    if mode == "markdown":
        monkeypatch.setattr(
            "converter.pipeline.write_office",
            lambda *args: pytest.fail("MD-only invoked Office writer"),
        )
    result = pipeline.run(job, provider, mode, tmp_path / "output")
    assert {
        Path(value["path"]).suffix for value in result.values() if value["status"] == "done"
    } == extensions
    assert all(value["status"] in {"done", "not_requested"} for value in result.values())
    original_calls = provider.calls
    doc = store.load_document(job)
    store.revise(job, {doc.segments[0].id: "Reviewed heading"})
    result = pipeline.run(job, None, mode, tmp_path / "output", export_only=True)
    assert provider.calls == original_calls
    assert all(
        value.get("revision", doc.revision + 1) == doc.revision + 1 for value in result.values()
    )
    assert file_hash(word_file) == doc.source_hash


def test_change_output_does_not_translate_again(word_file, store, provider, tmp_path):
    job = store.add(word_file)
    pipeline = Pipeline(store)
    pipeline.run(job, provider, "office", tmp_path / "output")
    calls = provider.calls
    pipeline.run(job, provider, "both", tmp_path / "output")
    assert provider.calls == calls
    assert store.get(job)["exports"]["markdown"]["status"] == "done"


def test_pdf_text_and_table_export(text_pdf, tmp_path):
    from docx import Document

    doc = read_pdf(text_pdf, tmp_path / "assets", language="eng")
    assert {"Document text", "Second paragraph", "Name", "Value", "Item", "100"} <= {
        s.source for s in doc.segments
    }
    assert len([s for s in doc.segments if s.source == "Name"]) == 1
    for segment in doc.segments:
        segment.target = "English " + segment.source
    output = tmp_path / "from-pdf.docx"
    write_pdf_docx(doc, output)
    word = Document(output)
    assert len(word.tables) == 1
    assert word.tables[0].cell(0, 0).text == "English Name"
    assert "English Document text" in render_markdown(doc, "assets")


def test_scanned_pdf_runs_real_ocr(tmp_path):
    import shutil

    from PIL import Image, ImageDraw, ImageFont
    from reportlab.pdfgen import canvas

    if not shutil.which("tesseract"):
        pytest.skip("System OCR not installed; covered by Windows bundled OCR smoke check")
    # Scanned image document: no PDF text layer. Use a clear available system font.
    candidates = [
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
        Path("C:/Windows/Fonts/arial.ttf"),
    ]
    font_path = next((path for path in candidates if path.exists()), None)
    if not font_path:
        pytest.skip("No test font installed")
    image = Image.new("RGB", (1800, 1200), "white")
    draw = ImageDraw.Draw(image)
    draw.text(
        (150, 180), "SCANNED DOCUMENT", font=ImageFont.truetype(str(font_path), 70), fill="black"
    )
    draw.text(
        (150, 330), "Invoice total 12345", font=ImageFont.truetype(str(font_path), 60), fill="black"
    )
    image_path = tmp_path / "scan.png"
    image.save(image_path)
    source = tmp_path / "scan.pdf"
    pdf = canvas.Canvas(str(source), pagesize=(600, 400))
    pdf.drawImage(str(image_path), 0, 0, width=600, height=400)
    pdf.save()
    doc = read_pdf(source, tmp_path / "assets", language="eng")
    text = " ".join(s.source for s in doc.segments)
    assert "SCANNED DOCUMENT" in text
    assert "12345" in text
    assert not doc.assets  # Do not duplicate scan background as a translated image.


def test_pdf_two_columns_follow_reading_order(tmp_path):
    from reportlab.pdfgen import canvas

    source = tmp_path / "columns.pdf"
    pdf = canvas.Canvas(str(source))
    for index in range(3):
        pdf.drawString(50, 750 - index * 30, f"Left column {index}")
        pdf.drawString(350, 750 - index * 30, f"Right column {index}")
    pdf.save()
    doc = read_pdf(source, tmp_path / "assets", language="eng")
    for segment in doc.segments:
        segment.target = segment.source
    markdown = render_markdown(doc, "assets")
    assert markdown.index("Left column 2") < markdown.index("Right column 0")


@pytest.mark.parametrize("source_fixture", ["slides_file", "spreadsheet", "text_pdf"])
@pytest.mark.parametrize("mode", ["office", "markdown", "both"])
def test_all_document_formats_support_three_modes(
    source_fixture, mode, request, store, provider, tmp_path
):
    source = request.getfixturevalue(source_fixture)
    job = store.add(source)
    result = Pipeline(store).run(job, provider, mode, tmp_path / "output", ocr_language="eng")
    selected = {"office", "markdown"} if mode == "both" else {mode}
    for kind, value in result.items():
        assert value["status"] == ("done" if kind in selected else "not_requested"), value
        if kind in selected:
            assert Path(value["path"]).stat().st_size > 0
    calls = provider.calls
    Pipeline(store).run(job, None, mode, tmp_path / "output", export_only=True)
    assert provider.calls == calls


def test_word_tabs_and_breaks_are_not_relocated(tmp_path):
    from docx import Document

    source = tmp_path / "breaks.docx"
    word = Document()
    paragraph = word.add_paragraph()
    run = paragraph.add_run("第一")
    run.add_tab()
    paragraph.add_run("第二").add_break()
    paragraph.add_run("第三")
    word.save(source)
    doc = read_office(source, tmp_path / "assets")
    translations = {"第一": "FIRST", "第二": "SECOND", "第三": "THIRD"}
    assert set(translations) <= {segment.source for segment in doc.segments}
    for segment in doc.segments:
        segment.target = translations[segment.source]
    output = tmp_path / "result.docx"
    write_office(doc, output)
    assert Document(output).paragraphs[0].text == "FIRST\tSECOND\nTHIRD"
