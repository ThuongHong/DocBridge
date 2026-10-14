import pytest
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfgen import canvas

from converter.markdown import render_markdown
from converter.models import Block, Document, Segment
from converter.pdf import PDF_ADAPTER_VERSION, read_pdf, write_pdf_docx
from converter.pipeline import Pipeline


def test_book_lines_form_paragraphs_and_keep_chinese_words_and_context(tmp_path):
    source = tmp_path / "book.pdf"
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    pdf = canvas.Canvas(str(source), pagesize=(420, 595))
    pdf.setFont("STSong-Light", 20)
    pdf.drawString(48, 550, "汽车电子系统")
    pdf.setFont("STSong-Light", 14)
    pdf.drawString(48, 500, "防抱死制动系统可以防止")
    pdf.drawString(48, 484, "车轮抱死。")
    pdf.drawString(48, 440, "请检查控制器的工作状态。")
    pdf.save()
    doc = read_pdf(source, tmp_path / "assets")
    sources = [segment.source for segment in doc.segments]
    assert "防抱死制动系统可以防止车轮抱死。" in sources
    assert "请检查控制器的工作状态。" in sources
    paragraph = next(s for s in doc.segments if "车轮抱死" in s.source)
    assert len(paragraph.anchor["line_bboxes"]) == 2
    assert "汽车电子系统" in paragraph.context
    assert "请检查控制器" in paragraph.context
    assert doc.adapter_version == PDF_ADAPTER_VERSION
    markdown = render_markdown(doc, "assets", source=True)
    assert "## 汽车电子系统" in markdown
    assert "防抱死制动系统可以防止车轮抱死。" in markdown


def test_pdf_docx_retains_page_sizes_fonts_image_size_and_real_content(
    tmp_path, image, monkeypatch
):
    from docx import Document
    from docx.enum.section import WD_ORIENT

    monkeypatch.setattr("converter.pdf.ocr_image", lambda *args: [])
    source = tmp_path / "layout.pdf"
    pdf = canvas.Canvas(str(source), pagesize=(420, 595))
    pdf.setFillColorRGB(0.2, 0.3, 0.4)
    pdf.setFont("Helvetica-Bold", 20)
    pdf.drawString(48, 545, "Technical manual")
    pdf.setFillColorRGB(0, 0, 0)
    pdf.setFont("Helvetica", 12)
    pdf.drawString(48, 505, "Voltage must be 12 V.")
    pdf.drawImage(str(image), 48, 400, width=100, height=40)
    pdf.showPage()
    pdf.setPageSize((595, 420))
    pdf.setFont("Times-Italic", 11)
    pdf.drawString(42, 360, "Pressure must be 100 kPa.")
    pdf.save()
    doc = read_pdf(source, tmp_path / "assets", language="eng")
    for segment in doc.segments:
        segment.target = segment.source
    output = tmp_path / "layout.docx"
    write_pdf_docx(doc, output)
    word = Document(output)
    assert [(round(s.page_width.pt), round(s.page_height.pt)) for s in word.sections] == [
        (420, 595), (595, 420)
    ]
    assert word.sections[1].orientation == WD_ORIENT.LANDSCAPE
    assert all(p.text not in {"Page 1", "Page 2"} for p in word.paragraphs)
    heading = next(p for p in word.paragraphs if p.text == "Technical manual")
    assert heading.runs[0].bold
    assert heading.runs[0].font.size.pt == 20
    assert str(heading.runs[0].font.color.rgb) == "334C66"
    italic = next(p for p in word.paragraphs if "Pressure" in p.text)
    assert italic.runs[0].italic
    assert italic.runs[0].font.size.pt == 11
    assert word.inline_shapes[0].width.pt == 100
    assert word.inline_shapes[0].height.pt == 40


def test_tightly_spaced_columns_do_not_merge_into_each_other(tmp_path):
    source = tmp_path / "columns.pdf"
    pdf = canvas.Canvas(str(source), pagesize=(612, 792))
    for index in range(3):
        pdf.drawString(48, 700 - index * 14, f"Left column line {index}")
        pdf.drawString(350, 700 - index * 14, f"Right column line {index}")
    pdf.save()
    doc = read_pdf(source, tmp_path / "assets")
    assert all(not ("Left" in s.source and "Right" in s.source) for s in doc.segments)
    markdown = render_markdown(doc, "assets", source=True)
    assert markdown.index("Left column line 2") < markdown.index("Right column line 0")
    for segment in doc.segments:
        segment.target = segment.source
    output = tmp_path / "columns.docx"
    write_pdf_docx(doc, output)
    from docx import Document as WordDocument
    word = WordDocument(output)
    assert len(word.tables) == 1
    table = word.tables[0]
    assert "Left column line 2" in table.cell(0, 0).text
    assert not table.cell(0, 1).text
    assert "Right column line 0" in table.cell(0, 2).text
    assert all("Right" not in p.text for p in table.cell(0, 0).paragraphs)
    assert all("Left" not in p.text for p in table.cell(0, 2).paragraphs)


def test_chinese_manual_columns_with_narrow_gutter_remain_separate(tmp_path):
    source = tmp_path / "narrow-gutter.pdf"
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    pdf = canvas.Canvas(str(source), pagesize=(595, 842))
    pdf.setFont("STSong-Light", 9)
    for index, label in enumerate(["第一", "第二", "第三"]):
        pdf.drawString(50, 750 - index * 14, f"左栏{label}段落" + "说明" * 10)
        pdf.drawString(304, 750 - index * 14, f"右栏{label}段落" + "说明" * 10)
    pdf.save()
    doc = read_pdf(source, tmp_path / "assets")
    assert all(not ("左栏" in s.source and "右栏" in s.source) for s in doc.segments)
    markdown = render_markdown(doc, "assets", source=True)
    assert markdown.index("左栏第三") < markdown.index("右栏第一")


def test_overprinted_native_text_is_extracted_once(tmp_path):
    source = tmp_path / "overprinted.pdf"
    pdf = canvas.Canvas(str(source))
    for _ in range(2):
        pdf.drawString(50, 750, "CAN ID 0x123 must remain unchanged.")
    pdf.save()
    doc = read_pdf(source, tmp_path / "assets")
    assert [s.source for s in doc.segments] == ["CAN ID 0x123 must remain unchanged."]


def test_parameter_table_without_body_column_rules_keeps_min_typ_max(tmp_path):
    source = tmp_path / "parameter-table.pdf"
    pdf = canvas.Canvas(str(source))
    edges = [40, 240, 320, 400, 480, 560]
    for y in [720, 700, 660]:
        pdf.line(edges[0], y, edges[-1], y)
    for x in edges:
        pdf.line(x, 700, x, 720)
    for x in [40, 240, 480, 560]:
        pdf.line(x, 660, x, 700)
    for left, right, text in zip(edges[:-1], edges[1:], ["Parameter", "MIN", "TYP", "MAX", "Unit"], strict=True):
        pdf.drawCentredString((left + right) / 2, 706, text)
    for left, right, text in zip(edges[:-1], edges[1:], ["VCM", "-12", "0", "12", "V"], strict=True):
        pdf.drawCentredString((left + right) / 2, 676, text)
    pdf.save()
    doc = read_pdf(source, tmp_path / "assets")
    table = next(b for b in doc.blocks[0].children if b.kind == "table")
    lookup = doc.by_id()
    assert ["".join(lookup[sid].source for sid in cell.segments)
            for cell in table.children[1].children] == ["VCM", "-12", "0", "12", "V"]


def test_repeated_margin_text_is_retained_but_excluded_from_body_context(tmp_path):
    source = tmp_path / "repeated-margins.pdf"
    pdf = canvas.Canvas(str(source), pagesize=(420, 595))
    for page in range(2):
        pdf.drawString(50, 575, "Technical manual")
        pdf.drawString(50, 450, f"Body text {page + 1}")
        pdf.drawString(50, 15, f"Copyright manufacturer - Page {page + 1}")
        pdf.showPage()
    pdf.save()
    doc = read_pdf(source, tmp_path / "assets")
    assert any(s.source.startswith("Copyright manufacturer") for s in doc.segments)
    assert any(b.meta.get("role") == "page_footer" for b in doc.blocks[0].children)
    body = next(s for s in doc.segments if s.source == "Body text 1")
    assert "Copyright" not in body.context
    markdown = render_markdown(doc, "assets", source=True)
    assert "Body text 1" in markdown and "Body text 2" in markdown
    assert "Copyright" not in markdown


@pytest.mark.parametrize("values", [(12, 24), (12, 12)])
def test_bottom_technical_values_are_never_hidden_as_page_furniture(tmp_path, values):
    source = tmp_path / "bottom-specifications.pdf"
    pdf = canvas.Canvas(str(source), pagesize=(420, 595))
    for value in values:
        pdf.drawString(50, 480, "Electrical specifications")
        pdf.drawString(50, 25, f"Maximum voltage: {value} V")
        pdf.showPage()
    pdf.save()
    doc = read_pdf(source, tmp_path / "assets")
    markdown = render_markdown(doc, "assets", source=True)
    for value in set(values):
        assert f"Maximum voltage: {value} V" in markdown
    assert markdown.count("Maximum voltage:") == 2


def test_open_table_cells_keep_native_text_inside_table_region(tmp_path):
    source = tmp_path / "open-cells.pdf"
    pdf = canvas.Canvas(str(source))
    for y in [720, 700, 680, 660]:
        pdf.line(40, y, 400, y)
    for x in [40, 160, 280, 400]:
        pdf.line(x, 700, x, 720)
    for x in [160, 280]:
        pdf.line(x, 660, x, 700)
    for x, label in [(50, "ID"), (170, "Meaning"), (290, "Purpose")]:
        pdf.drawString(x, 706, label)
    for y, row in [(686, ["CAN_ID", "Bus message", "Diagnostics"]),
                   (666, ["DTC", "Fault code", "Repair"])]:
        for x, value in zip([50, 170, 290], row, strict=True):
            pdf.drawString(x, y, value)
    pdf.save()
    doc = read_pdf(source, tmp_path / "assets")
    table = next(b for b in doc.blocks[0].children if b.kind == "table")
    lookup = doc.by_id()
    assert ["".join(lookup[sid].source for sid in c.segments)
            for c in table.children[1].children] == ["CAN_ID", "Bus message", "Diagnostics"]


def test_picture_ocr_does_not_interrupt_a_wrapped_native_sentence():
    from converter.pdf import reading_order
    from converter.pdf_layout import merge_paragraph_lines

    doc = Document("source.pdf", "hash", "pdf")
    items = [(50, 40, Block("image", meta={"asset": "figure.png", "bbox": [40, 50, 310, 140]}))]
    for text, box, meta, style in [
        ("A sentence continues", [40, 100, 300, 110], {}, {"font_size": 12}),
        ("OCR in picture", [40, 105, 120, 115],
         {"role": "image_text", "image_asset": "figure.png"}, {}),
        ("on the next line.", [40, 114, 130, 124], {}, {"font_size": 12}),
    ]:
        segment = doc.add_segment(text, {"bbox": box}, style=style)
        block = Block("paragraph", [segment.id], meta={"bbox": box} | meta)
        items.append((box[1], box[0], block))
    ordered = reading_order(items, doc, 600)
    merge_paragraph_lines(ordered, doc, 600)
    assert any(s.source == "A sentence continues on the next line." for s in doc.segments)
    assert any(s.source == "OCR in picture" for s in doc.segments)


def test_numbered_instruction_keeps_hanging_continuation_in_one_segment(tmp_path):
    source = tmp_path / "instruction.pdf"
    pdf = canvas.Canvas(str(source))
    pdf.setFont("Helvetica", 9)
    pdf.drawString(50, 750, "3. Read the parameter name and result")
    pdf.drawString(62.4, 737.6, "and unit.")
    pdf.drawString(50, 712, "4. Exit the diagnostic function.")
    pdf.save()
    doc = read_pdf(source, tmp_path / "assets")
    assert any(s.source == "3. Read the parameter name and result and unit." for s in doc.segments)
    assert any(s.source == "4. Exit the diagnostic function." for s in doc.segments)


def test_full_width_heading_and_footer_bound_column_bands(tmp_path):
    from docx import Document as WordDocument
    source = tmp_path / "bands.pdf"
    pdf = canvas.Canvas(str(source), pagesize=(612, 792))
    pdf.setFont("Helvetica-Bold", 20)
    pdf.drawString(48, 740, "Full width technical document heading")
    pdf.setFont("Helvetica", 12)
    for index in range(3):
        pdf.drawString(48, 700 - index * 14, f"Left column paragraph {index}")
        pdf.drawString(350, 700 - index * 14, f"Right column paragraph {index}")
    pdf.drawString(48, 600, "Full width concluding note across the two document columns.")
    pdf.save()
    doc = read_pdf(source, tmp_path / "assets")
    for segment in doc.segments:
        segment.target = segment.source
    output = tmp_path / "bands.docx"
    write_pdf_docx(doc, output)
    word = WordDocument(output)
    assert word.paragraphs[0].text == "Full width technical document heading"
    assert word.paragraphs[-1].text.startswith("Full width concluding note")
    assert len(word.tables) == 1
    assert "Left column" in word.tables[0].cell(0, 0).text
    assert "Right column" in word.tables[0].cell(0, 2).text


def test_merging_on_first_page_keeps_ids_and_translations_separate(tmp_path, store):
    source = tmp_path / "multipage.pdf"
    pdf = canvas.Canvas(str(source), pagesize=(420, 595))
    pdf.drawString(48, 500, "First paragraph continues")
    pdf.drawString(48, 486, "onto another line.")
    pdf.drawString(48, 420, "Figure caption")
    pdf.showPage()
    pdf.drawString(48, 500, "Second page title")
    pdf.save()

    class ReferenceProvider:
        batch_chars = 6000
        handles_long_segments = True
        def fingerprint(self, glossary):
            return "layout-reference"
        def translate(self, segments, glossary):
            return {s.id: "Verified " + s.source for s in segments}

    job = store.add(source)
    Pipeline(store).run(job, ReferenceProvider(), "markdown", tmp_path / "out", ocr_language="eng")
    doc = store.load_document(job)
    assert len({s.id for s in doc.segments}) == len(doc.segments)
    assert len(doc.segments) == 3
    assert all(s.target == "Verified " + s.source for s in doc.segments)


def test_duplicate_ids_fail_before_content_can_be_mapped_to_the_wrong_segment():
    doc = Document("source.pdf", "hash", "pdf", segments=[
        Segment("s000000", "First page"), Segment("s000000", "Second page"),
    ])
    with pytest.raises(ValueError, match="IDs"):
        doc.by_id()


def test_ocr_paragraph_ids_prevent_merging_separate_paragraphs(tmp_path, image, monkeypatch):
    source = tmp_path / "scan.pdf"
    pdf = canvas.Canvas(str(source), pagesize=(420, 595))
    pdf.drawImage(str(image), 0, 0, width=420, height=595)
    pdf.save()
    monkeypatch.setattr("converter.pdf.configure_tesseract", lambda: "")
    monkeypatch.setattr("converter.pdf.ocr_image", lambda *args: [
        {"text": "防抱死制动系统可以防止", "bbox": [144, 180, 1200, 222],
         "confidence": 90, "paragraph": [1, 1]},
        {"text": "车轮抱死。", "bbox": [144, 228, 1200, 270],
         "confidence": 55, "paragraph": [1, 1]},
        {"text": "请检查控制器。", "bbox": [144, 276, 1200, 318],
         "confidence": 95, "paragraph": [1, 2]},
    ])
    doc = read_pdf(source, tmp_path / "assets")
    assert len(doc.segments) == 2
    assert doc.segments[0].source == "防抱死制动系统可以防止车轮抱死。"
    assert "55%" in doc.segments[0].warning
    assert doc.segments[1].source == "请检查控制器。"


@pytest.mark.parametrize("older_version", [1, 2, 3, 4])
def test_changed_pdf_adapter_rejects_resume_but_keeps_existing_exports(text_pdf, store, tmp_path, older_version):
    job = store.add(text_pdf)
    pipeline = Pipeline(store)
    pipeline.run(job, None, "markdown", tmp_path / "out", operation="convert", ocr_language="eng")
    doc = store.load_document(job)
    doc.adapter_version = older_version
    store.save_document(job, doc)
    with pytest.raises(ValueError, match="adapter"):
        pipeline.parse(job, "eng")
    result = pipeline.run(
        job, None, "markdown", tmp_path / "out", operation="convert", export_only=True
    )
    assert result["markdown"]["status"] == "done"
