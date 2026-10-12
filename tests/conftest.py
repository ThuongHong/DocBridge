import os

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import pytest
from PIL import Image

from converter.store import Store


class FixtureProvider:
    batch_chars = 150

    def __init__(self):
        self.calls = 0

    def fingerprint(self, glossary):
        return "fixture-profile"

    def translate(self, segments, glossary):
        self.calls += 1
        return {s.id: "Translated " + s.source for s in segments}


@pytest.fixture
def provider():
    return FixtureProvider()


@pytest.fixture
def store(tmp_path):
    return Store(tmp_path / "store")


@pytest.fixture
def image(tmp_path):
    path = tmp_path / "picture.png"
    Image.new("RGB", (40, 40), color="blue").save(path)
    return path


@pytest.fixture
def word_file(tmp_path, image):
    from docx import Document
    from docx.shared import Inches

    doc = Document()
    doc.add_heading("中文标题", level=1)
    paragraph = doc.add_paragraph()
    paragraph.add_run("重要 ").bold = True
    paragraph.add_run("普通").italic = True
    # Artificial run split with the same style must not fragment the translation.
    p = doc.add_paragraph()
    p.add_run("测试")
    p.add_run("文本")
    table = doc.add_table(rows=2, cols=2)
    for i, text in enumerate(["名称", "数值", "商品", "100"]):
        table.cell(i // 2, i % 2).text = text
    doc.sections[0].header.paragraphs[0].text = "页眉"
    doc.sections[0].footer.paragraphs[0].text = "页脚"
    doc.add_picture(str(image), width=Inches(1))
    path = tmp_path / "文档.docx"
    doc.save(path)
    return path


@pytest.fixture
def slides_file(tmp_path, image):
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    box = group.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    p = box.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = "粗体"
    run.font.bold = True
    run = p.add_run()
    run.text = "文本"
    run.font.italic = True
    slide.shapes.add_picture(str(image), Inches(1), Inches(3), width=Inches(1))
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(5), Inches(4), Inches(1)).table
    table.cell(0, 0).text = "列名"
    table.cell(1, 1).text = "数据"
    slide.notes_slide.notes_text_frame.text = "演讲备注"
    path = tmp_path / "slides.pptx"
    presentation.save(path)
    return path


@pytest.fixture
def spreadsheet(tmp_path, image):
    import xlsxwriter

    path = tmp_path / "sheets.xlsx"
    workbook = xlsxwriter.Workbook(path)
    sheet = workbook.add_worksheet("原始名称")
    sheet.write("A1", "相同")
    sheet.write("A2", "相同")
    sheet.write("B1", 123)
    sheet.write_formula("B2", "=B1*2", None, 246)
    bold = workbook.add_format({"bold": True})
    italic = workbook.add_format({"italic": True})
    sheet.write_rich_string("C1", bold, "粗体", italic, "斜体")
    sheet.insert_image("E1", str(image))
    chart = workbook.add_chart({"type": "column"})
    chart.add_series({"values": "='原始名称'!$B$1:$B$2"})
    sheet.insert_chart("E6", chart)
    workbook.close()
    return path


@pytest.fixture
def text_pdf(tmp_path):
    from reportlab.pdfgen import canvas

    path = tmp_path / "document.pdf"
    pdf = canvas.Canvas(str(path))
    pdf.drawString(50, 780, "Document text")
    pdf.drawString(50, 750, "Second paragraph")
    # Ruled table: two rows, two columns.
    for y in [600, 630, 660]:
        pdf.line(50, y, 250, y)
    for x in [50, 150, 250]:
        pdf.line(x, 600, x, 660)
    pdf.drawString(60, 640, "Name")
    pdf.drawString(160, 640, "Value")
    pdf.drawString(60, 610, "Item")
    pdf.drawString(160, 610, "100")
    pdf.save()
    return path
