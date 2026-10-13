from PIL import Image, ImageDraw
from reportlab.pdfgen import canvas

from converter.pdf import read_pdf, write_pdf_docx
from converter.scan_layout import remaining_figures, ruled_tables


def test_weak_line_retry_retains_original_geometry_and_rejects_lower_confidence(monkeypatch):
    from converter.pdf import refine_ocr_lines
    image = Image.new("RGB", (400, 200), "white")
    source = {"text":"汐查控制器和传感器之间的连接。", "bbox":[20,70,370,105],
              "confidence":76, "minimum_confidence":0, "paragraph":[2,1], "text_height":33}
    calls = []
    def recognize(crop, language, config, psm):
        calls.append((crop.size, psm))
        return [{"text":"检查控制器和传感器之间的连接。", "bbox":[0,10,350,43],
                 "confidence":94, "minimum_confidence":90, "paragraph":[1,1], "text_height":33}]
    monkeypatch.setattr("converter.pdf.ocr_image", recognize)
    improved = refine_ocr_lines(image, [source], "chi_sim+eng", "")
    assert improved[0]["text"] == "检查控制器和传感器之间的连接。"
    assert improved[0]["bbox"] == source["bbox"]
    assert improved[0]["paragraph"] == source["paragraph"]
    assert calls[0][1] == 7
    strong = source | {"minimum_confidence":90, "confidence":94}
    assert refine_ocr_lines(image,[strong],"chi_sim+eng","") == [strong]
    assert len(calls) == 1
    monkeypatch.setattr("converter.pdf.ocr_image", lambda *args, **kwargs: [source | {"confidence":60}])
    assert refine_ocr_lines(image,[source],"chi_sim+eng","") == [source]


def test_fragmented_line_stroke_retry_keeps_numbers_and_original_coordinates(monkeypatch):
    from converter.pdf import refine_ocr_lines
    image = Image.new("RGB", (500,200), "white")
    source = {"text":"在电压低于12 V时，不得局用该功能。", "bbox":[20,70,450,105],
              "confidence":89, "minimum_confidence":73, "paragraph":[2,1], "text_height":33}
    responses = iter([[source | {"confidence":88, "minimum_confidence":63}],
                      [source | {"text":"在电压低于12 V时，不得启用该功能。",
                                 "confidence":91, "minimum_confidence":75}]])
    monkeypatch.setattr("converter.pdf.ocr_image", lambda *args, **kwargs:next(responses))
    result = refine_ocr_lines(image,[source],"chi_sim+eng","",fragmented=True)[0]
    assert "启用" in result["text"] and "12 V" in result["text"]
    assert result["bbox"] == source["bbox"] and result["paragraph"] == source["paragraph"]
    monkeypatch.setattr("converter.pdf.ocr_image", lambda *args, **kwargs:[
        source | {"text":"在电压低于21 V时，不得启用该功能。", "confidence":99, "minimum_confidence":99}])
    assert refine_ocr_lines(image,[source],"chi_sim+eng","",fragmented=True) == [source]


def scan_image():
    image = Image.new("RGB", (420, 595), "white")
    draw = ImageDraw.Draw(image)
    for y in (235, 265, 295, 325):
        draw.line((48, y, 372, y), fill="black", width=1)
    for x in (48, 230, 372):
        draw.line((x, 235, x, 325), fill="black", width=1)
    draw.rectangle((48, 400, 88, 440), outline="#2463a6")
    draw.rectangle((108, 400, 148, 440), outline="#2463a6")
    draw.line((88, 420, 108, 420), fill="#2463a6")
    return image


def test_grid_detection_excludes_diagrams_and_residual_preserves_their_pixels():
    image = scan_image()
    grids = ruled_tables(image)
    assert len(grids) == 1
    assert grids[0]["xs"] == [48, 230, 372]
    assert grids[0]["ys"] == [235, 265, 295, 325]
    figures = remaining_figures(image, [], [grids[0]["bbox"]], 12)
    assert figures == [[48, 400, 149, 441]]


def test_antialiased_corner_does_not_hide_an_otherwise_complete_grid():
    image = scan_image()
    ImageDraw.Draw(image).point((372, 235), fill="white")
    assert len(ruled_tables(image)) == 1


def test_scanned_table_cells_and_diagram_survive_pdf_export(tmp_path, monkeypatch):
    from docx import Document
    image = scan_image()
    image_path = tmp_path / "scan.png"
    image.save(image_path)
    source = tmp_path / "scan.pdf"
    pdf = canvas.Canvas(str(source), pagesize=(420, 595))
    pdf.drawImage(str(image_path), 0, 0, width=420, height=595)
    pdf.save()
    cells = iter(["参数", "数值", "工作电压", "12 V", "压力", "100 kPa"])

    def recognize(image, language, config, psm=3):
        if psm == 3:
            return [{"text": "100 kPa", "bbox": [720, 900, 900, 942],
                     "confidence": 95, "paragraph": [1, 1]}]
        return [{"text": next(cells), "bbox": [5, 5, 30, 38], "text_height": 33,
                 "confidence": 95, "paragraph": [1, 1]}]

    monkeypatch.setattr("converter.pdf.configure_tesseract", lambda: "")
    monkeypatch.setattr("converter.pdf.ocr_image", recognize)
    doc = read_pdf(source, tmp_path / "assets")
    assert [s.source for s in doc.segments] == ["参数", "数值", "工作电压", "12 V", "压力", "100 kPa"]
    assert [b.kind for b in doc.blocks[0].children] == ["table", "image"]
    assert len(doc.assets) == 1
    for segment in doc.segments:
        segment.target = segment.source
    output = tmp_path / "scan.docx"
    write_pdf_docx(doc, output)
    word = Document(output)
    assert len(word.tables) == 1 and len(word.inline_shapes) == 1
    assert word.tables[0].cell(1, 1).text == "12 V"
    assert word.tables[0].cell(2, 1).text == "100 kPa"
    assert 99 <= word.inline_shapes[0].width.pt <= 102


def test_unread_scanned_cell_keeps_table_snapshot_instead_of_losing_content(tmp_path, monkeypatch):
    image = scan_image()
    ImageDraw.Draw(image).rectangle((65, 245, 90, 255), fill="black")
    image_path = tmp_path / "scan.png"
    image.save(image_path)
    source = tmp_path / "unread.pdf"
    pdf = canvas.Canvas(str(source), pagesize=(420, 595))
    pdf.drawImage(str(image_path), 0, 0, width=420, height=595)
    pdf.save()
    monkeypatch.setattr("converter.pdf.configure_tesseract", lambda: "")
    monkeypatch.setattr("converter.pdf.ocr_image", lambda *args: [])
    doc = read_pdf(source, tmp_path / "assets")
    assert all(b.kind != "table" for b in doc.blocks[0].children)
    assert len(doc.assets) == 2
    assert any("bảng scan chưa đọc đủ" in warning for warning in doc.warnings)


def test_fragmented_ocr_retries_sparse_mode_only_when_confidence_improves(tmp_path, monkeypatch):
    image_path = tmp_path / "page.png"
    Image.new("RGB", (420, 595), "white").save(image_path)
    source = tmp_path / "sparse.pdf"
    pdf = canvas.Canvas(str(source), pagesize=(420, 595))
    pdf.drawImage(str(image_path), 0, 0, width=420, height=595)
    pdf.save()
    calls = []

    def recognize(image, language, config, psm=3):
        calls.append(psm)
        if psm == 11:
            return [{"text": "在电压低于12 V时，不得启用该功能。", "bbox": [144, 200, 1000, 245],
                     "confidence": 91, "paragraph": [1, 1]}]
        return [{"text": text, "bbox": [144 + index * 70, 200, 204 + index * 70, 245],
                 "confidence": confidence, "paragraph": [index + 1, 1]}
                for index, (text, confidence) in enumerate([
                    ("在日", 73), ("壓低", 29), ("F12V 时，不得局/", 85), ("該功", 89), ("HEo", 86)])]

    monkeypatch.setattr("converter.pdf.configure_tesseract", lambda: "")
    monkeypatch.setattr("converter.pdf.ocr_image", recognize)
    doc = read_pdf(source, tmp_path / "assets")
    assert calls == [3, 11]
    assert [s.source for s in doc.segments] == ["在电压低于12 V时，不得启用该功能。"]
