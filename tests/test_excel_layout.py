import zipfile

import openpyxl
import pytest
from openpyxl.styles import Alignment, Font

from converter.office import read_office, write_office

TEXT = 'The construction safety measures must be inspected before work begins.'


def translated(source, tmp_path, targets=None):
    doc = read_office(source, tmp_path / 'assets')
    for segment in doc.segments:
        segment.target = (targets or {}).get(segment.anchor.get('cell'), TEXT)
    return doc


def test_readable_wraps_and_grows_changed_cells_without_mutating_shared_styles(tmp_path):
    source = tmp_path / 'source.xlsx'
    book = openpyxl.Workbook()
    sheet = book.active
    for address in ['A1', 'A2']:
        sheet[address] = '控制器'
        sheet[address].font = Font(name='Arial', size=11)
        sheet[address].alignment = Alignment(wrap_text=False, vertical='top')
    sheet.column_dimensions['A'].width = 10
    sheet.row_dimensions[1].height = 16
    sheet['B1'] = '=1+2'
    book.save(source)
    doc = translated(source, tmp_path, {'A2': '控制器'})
    preserve, readable = tmp_path / 'preserve.xlsx', tmp_path / 'readable.xlsx'
    write_office(doc, preserve)
    write_office(doc, readable, excel_layout='readable')
    original = openpyxl.load_workbook(preserve).active
    result = openpyxl.load_workbook(readable).active
    assert original.row_dimensions[1].height == 16
    assert result.row_dimensions[1].height > 16
    assert result['A1'].alignment.wrap_text
    assert not result['A2'].alignment.wrap_text
    assert result['A2'].value == '控制器'
    assert result['B1'].value == '=1+2'
    assert result.column_dimensions['A'].width == 10
    with zipfile.ZipFile(source) as before, zipfile.ZipFile(preserve) as after:
        assert before.read('xl/styles.xml') == after.read('xl/styles.xml')
    assert doc.coverage['layout']['workbooks'][0]['increased_rows'] == 1
    assert 'han_locations' in doc.coverage


def test_merged_widths_vertical_height_distribution_newlines_and_hidden_rows(tmp_path):
    source = tmp_path / 'merged.xlsx'
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.column_dimensions['A'].width = sheet.column_dimensions['B'].width = 10
    sheet.merge_cells('A1:B1')
    sheet.merge_cells('A5:A7')
    for address in ['A1', 'A3', 'A5', 'A9']:
        sheet[address] = '检查'
        sheet[address].font = Font(name='Arial', size=11)
    for row in [1, 3, 5, 6, 7, 9]:
        sheet.row_dimensions[row].height = 16
    sheet.row_dimensions[6].hidden = True
    book.save(source)
    doc = translated(source, tmp_path, {'A9': 'One\nTwo\nThree'})
    output = tmp_path / 'out.xlsx'
    write_office(doc, output, excel_layout='readable')
    result = openpyxl.load_workbook(output).active
    assert result.row_dimensions[1].height < result.row_dimensions[3].height
    assert result.row_dimensions[5].height > 16 and result.row_dimensions[7].height > 16
    assert result.row_dimensions[6].hidden and result.row_dimensions[6].height == 16
    assert result.row_dimensions[9].height > 30
    assert set(map(str, result.merged_cells.ranges)) == {'A1:B1', 'A5:A7'}


def test_layout_respects_caps_rotation_and_hidden_columns(tmp_path):
    source = tmp_path / 'caps.xlsx'
    book = openpyxl.Workbook()
    sheet = book.active
    sheet['A1'] = sheet['A2'] = sheet['A3'] = '检查'
    sheet['A2'].alignment = Alignment(text_rotation=45)
    sheet.row_dimensions[3].hidden = True
    sheet.column_dimensions['A'].width = 2
    book.save(source)
    doc = translated(source, tmp_path, {'A1': 'Longword' * 500})
    output = tmp_path / 'out.xlsx'
    write_office(doc, output, excel_layout='readable')
    result = openpyxl.load_workbook(output).active
    assert result.row_dimensions[1].height == 409
    assert result['A2'].alignment.text_rotation == 45
    assert result.row_dimensions[3].hidden
    issues = doc.coverage['layout']['workbooks'][0]['issues']
    assert issues['row_height_cap_unresolved']['count'] == 1
    assert issues['rotated_text_not_fitted']['count'] == 1
    assert issues['hidden_cells_skipped']['count'] == 1


def test_inherited_row_and_column_font_styles_are_measured(tmp_path):
    source = tmp_path / 'inherited.xlsx'
    book = openpyxl.Workbook()
    sheet = book.active
    for address in ['A1', 'A2', 'B3']:
        sheet[address] = '检查'
    sheet.row_dimensions[1].font = Font(name='Arial', size=24)
    sheet.column_dimensions['B'].font = Font(name='Arial', size=28)
    sheet.column_dimensions['A'].width = sheet.column_dimensions['B'].width = 10
    book.save(source)
    doc = translated(source, tmp_path, {'A1': 'Check safety', 'A2': 'Check safety', 'B3': 'Check safety'})
    output = tmp_path / 'out.xlsx'
    write_office(doc, output, excel_layout='readable')
    result = openpyxl.load_workbook(output).active
    assert result['A1'].font.sz == 24
    assert result['B3'].font.sz == 28
    assert result.row_dimensions[1].height > result.row_dimensions[2].height
    assert result.row_dimensions[3].height > result.row_dimensions[2].height


def test_default_hidden_absent_merged_rows_do_not_absorb_visible_height(tmp_path):
    source = tmp_path / 'default-hidden.xlsx'
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.sheet_format.zeroHeight = True
    sheet['A1'] = '检查'
    sheet.merge_cells('A1:A3')
    sheet.row_dimensions[1].hidden = False
    sheet.column_dimensions['A'].width = 10
    book.save(source)
    # Explicit visible first row, with rows2/3 absent and hidden by default.
    from lxml import etree

    from converter.office import NS
    patched = tmp_path / 'patched.xlsx'
    with zipfile.ZipFile(source) as before, zipfile.ZipFile(patched, 'w') as after:
        for entry in before.infolist():
            content = before.read(entry.filename)
            if entry.filename == 'xl/worksheets/sheet1.xml':
                root = etree.fromstring(content)
                for row in root.findall('s:sheetData/s:row', NS):
                    if row.get('r') == '1':
                        row.set('hidden', '0')
                    else:
                        row.getparent().remove(row)
                content = etree.tostring(root)
            after.writestr(entry, content)
    doc = translated(patched, tmp_path)
    output = tmp_path / 'out.xlsx'
    write_office(doc, output, excel_layout='readable')
    result = openpyxl.load_workbook(output).active
    assert result.row_dimensions[1].height > 45
    assert not result.row_dimensions[1].hidden
    assert result.sheet_format.zeroHeight
    assert 2 not in result.row_dimensions and 3 not in result.row_dimensions


def test_invalid_layout_mode_rejected(tmp_path, spreadsheet):
    doc = translated(spreadsheet, tmp_path)
    with pytest.raises(ValueError, match='Excel layout'):
        write_office(doc, tmp_path / 'out.xlsx', excel_layout='invalid')


def test_justified_text_can_fit_and_capped_plain_text_scales_its_own_font(tmp_path):
    source = tmp_path / 'reduction.xlsx'
    book = openpyxl.Workbook()
    sheet = book.active
    for address in ['A1', 'A2']:
        sheet[address] = '检查'
        sheet[address].font = Font(name='Arial', size=16)
        sheet[address].alignment = Alignment(horizontal='justify')
    sheet.column_dimensions['A'].width = 60
    sheet.row_dimensions[1].height = 409
    book.save(source)
    doc = translated(source, tmp_path, {'A1': TEXT * 20, 'A2': '检查'})
    output = tmp_path / 'out.xlsx'
    write_office(doc, output, excel_layout='readable')
    result = openpyxl.load_workbook(output).active
    assert 9 <= result['A1'].font.sz < 16
    assert result['A2'].font.sz == 16
    assert result['A1'].alignment.horizontal == 'justify'
    assert result['A1'].value == TEXT * 20
    report = doc.coverage['layout']['workbooks'][0]
    assert report['reduced_font_cells'] == 1
    assert 'row_height_cap_unresolved' not in report['issues']


def test_repeated_cell_measurements_and_wrap_styles_are_reused(tmp_path, monkeypatch):
    import converter.excel_layout as layout

    source = tmp_path / 'repeat.xlsx'
    book = openpyxl.Workbook()
    for row in range(1, 1001):
        book.active.cell(row, 1, '检查')
    book.active.column_dimensions['A'].width = 10
    book.save(source)
    measured = layout.measured_height
    calls = []

    def counted(*args):
        calls.append(1)
        return measured(*args)

    monkeypatch.setattr(layout, 'measured_height', counted)
    doc = translated(source, tmp_path)
    output = tmp_path / 'out.xlsx'
    write_office(doc, output, excel_layout='readable')
    assert len(calls) == 1
    report = doc.coverage['layout']['workbooks'][0]
    assert report['wrapped_cells'] == report['increased_rows'] == 1000
    result = openpyxl.load_workbook(output)
    assert len(result._cell_styles) == 2


@pytest.mark.parametrize('columns', [1, 2])
def test_readable_fits_wide_print_area_to_one_page_width_without_height_limit(tmp_path, columns):
    from openpyxl.worksheet.pagebreak import Break

    source = tmp_path / 'print.xlsx'
    book = openpyxl.Workbook()
    sheet = book.active
    sheet['A1'] = '检查'
    sheet.column_dimensions['A'].width = 100.5
    if columns == 2:
        sheet['B1'] = 42
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.orientation = 'portrait'
    sheet.page_setup.scale = 73
    sheet.row_breaks.append(Break(id=20))
    book.save(source)
    doc = translated(source, tmp_path)
    output = tmp_path / 'out.xlsx'
    write_office(doc, output, excel_layout='readable')
    result = openpyxl.load_workbook(output).active
    assert result.page_setup.fitToWidth == 1 and result.page_setup.fitToHeight == 0
    if columns == 1:
        assert not result.sheet_properties.pageSetUpPr.fitToPage
        assert 10 <= result.page_setup.scale < 100
    else:
        assert result.sheet_properties.pageSetUpPr.fitToPage
        assert result.page_setup.scale is None
    assert result.page_setup.paperSize == int(sheet.PAPERSIZE_A4)
    assert result.page_setup.orientation == 'portrait'
    assert result.page_margins == sheet.page_margins
    assert result.row_breaks.brk[0].id == 20
    assert doc.coverage['layout']['workbooks'][0]['print_fit_sheets'] == 1
